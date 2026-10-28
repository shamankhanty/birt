import json
from datetime import date
from pathlib import Path
import sys
import tempfile
import unittest
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts/pipeline'))
from adapters import adapt_hospital_count

class HospitalCountTests(unittest.TestCase):
    def fixture(self,folder,missing=False,duplicate=False):
        p=Path(folder)
        (p/'operational-mo.json').write_text(json.dumps({'hospitalCount': {'name':'Количество госпитализаций','date':'20.10.2026','period':'01.01–20.10.2026','rows':[{'name':'MO','oid':'1.2.3','fact':80},{'name':'Absent','oid':'1.2.4','fact':5}]},'unrelated':{'rows':[{'fact':99}]}}))
        (p/'mo-data.json').write_text('{"hospital":{"numerator":75,"denominator":100}}')
        book=Workbook();ws=book.active;ws.title='v2'
        for row in ([6,7] if duplicate else [6]):
            ws.cell(row,2,'MO');ws.cell(row,3,'1.2.3')
            if not missing:ws.cell(row,4,100)
            ws.cell(row,5,70);ws.cell(row,6,5)
        source=p/'hospital.xlsx';book.save(source);return p,source
    def test_counts_use_column_d_preserve_missing_rows_and_leave_ratio_untouched(self):
        with tempfile.TemporaryDirectory() as folder:
            p,source=self.fixture(folder);before=(p/'mo-data.json').read_bytes()
            result=adapt_hospital_count(p,source,date(2026,10,25))
            data=json.loads((p/'operational-mo.json').read_text());ds=data['hospitalCount']
            self.assertEqual(ds['total'],100);self.assertEqual(ds['date'],'25.10.2026')
            self.assertEqual(ds['rows'][0]['fact'],100);self.assertEqual(ds['rows'][0]['previous'],80)
            self.assertEqual(ds['rows'][0]['trend'],20);self.assertIsNone(ds['rows'][1]['fact'])
            self.assertEqual(data['unrelated']['rows'][0]['fact'],99)
            self.assertEqual((p/'mo-data.json').read_bytes(),before)
            first=(p/'operational-mo.json').read_bytes();adapt_hospital_count(p,source,date(2026,10,25))
            self.assertEqual((p/'operational-mo.json').read_bytes(),first)
    def test_distinct_source_rows_with_same_oid_are_aggregated_without_losing_volume(self):
        with tempfile.TemporaryDirectory() as folder:
            p,source=self.fixture(folder,duplicate=True)
            from openpyxl import load_workbook
            book=load_workbook(source);book.active.cell(7,2,'MO Branch');book.active.cell(7,4,50);book.save(source)
            result=adapt_hospital_count(p,source,date(2026,10,25))
            data=json.loads((p/'operational-mo.json').read_text())['hospitalCount']
            self.assertEqual(data['total'],150)
            self.assertEqual(data['rows'][0]['fact'],150)
            self.assertEqual(data['rows'][0]['sourceRows'],2)
    def test_missing_column_d_is_not_replaced_with_zero(self):
        with tempfile.TemporaryDirectory() as folder:
            p,source=self.fixture(folder,missing=True)
            with self.assertRaisesRegex(ValueError,'столбце D'):adapt_hospital_count(p,source,date(2026,10,25))
    def test_duplicate_oid_does_not_double_the_case_total(self):
        with tempfile.TemporaryDirectory() as folder:
            p,source=self.fixture(folder,duplicate=True)
            with self.assertRaisesRegex(ValueError,'повтор идентичной строки'):adapt_hospital_count(p,source,date(2026,10,25))

if __name__=='__main__':unittest.main()
