import sys, tempfile, unittest, json
from pathlib import Path
from datetime import date
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts/pipeline'))
from adapters import adapt_tvsp_subunits

class MissingOids(unittest.TestCase):
    def test_only_mandatory_nontransmitting_oids_are_saved(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);app=root/'app';app.mkdir()
            for name in ['unit-data','mo-data','mo-details','mo-detail-oids']:
                (app/(name+'.json')).write_text('{}')
            prefix='1.2.643.5.1.13.13.12.2.16.1179.0.'
            wb=Workbook();ws=wb.active;ws.title='Детализация по СП'
            ws.cell(7,1,'Республика Татарстан');ws.cell(7,2,'ГП №21');ws.cell(7,3,prefix.rsplit('.0.',1)[0]);ws.cell(7,4,'101896')
            ws.cell(7,5,', '.join(prefix+x for x in ['1','2','3']))
            ws.cell(7,6,', '.join(prefix+x for x in ['1','999']))
            p=root/'source.xlsx';wb.save(p)
            adapt_tvsp_subunits(app,p,date(2026,10,29),'tvspAmbulatory','tvsp_ambulatory','ТВСП','объект')
            row=json.loads((app/'unit-data.json').read_text())['tvspAmbulatory']['rows'][0]
            self.assertEqual(row['missingUnitOids'],[prefix+'2',prefix+'3'])
            self.assertEqual(row['registeredSubunits'],1)
            self.assertEqual(row['plannedSubunits'],3)
if __name__=='__main__':unittest.main()
