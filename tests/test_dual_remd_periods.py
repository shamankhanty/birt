import sys, tempfile, unittest
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts/pipeline'))
from source_catalog import scan

class DualErrorPeriods(unittest.TestCase):
    def source(self,root,name,start):
        book=Workbook();ws=book.active
        ws.append([]);ws.append(['Дата начала периода',start]);ws.append(['Дата окончания периода',datetime(2026,10,28)])
        ws.append(['Дата формирования отчета:',datetime(2026,10,29)])
        ws.cell(5,1,'Отказы в регистрации ЭМД в РЭМД');ws.cell(7,2,'1.2.3');ws.cell(7,3,'МО');ws.cell(7,9,'Ошибка');ws.cell(7,10,7)
        book.save(root/name)
    def test_weekly_and_cumulative_same_cut_are_not_duplicates(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            self.source(root,'Отчёт - Отказы 01.01-28.10.2026.xlsx',datetime(2026,1,1))
            self.source(root,'Отчёт - Отказы 22.10-28.10.2026.xlsx',datetime(2026,10,22))
            report=scan(root)
            self.assertEqual(len(report['files']),2)
            self.assertTrue(all(x['family']=='remd_errors' and x['status']!='FAIL' for x in report['files']))
    def test_identical_period_still_blocks_duplicate_inputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            for suffix in ['a','b']:self.source(root,f'Отчёт - Отказы 01.01-28.10.2026 {suffix}.xlsx',datetime(2026,1,1))
            self.assertEqual(scan(root)['summary']['FAIL'],2)
if __name__=='__main__':unittest.main()
