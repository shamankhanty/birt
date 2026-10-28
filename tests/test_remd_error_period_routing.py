import json
from datetime import date
from pathlib import Path
import sys
import tempfile
import unittest
from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts/pipeline'))
from adapters import adapt_errors

class ErrorPeriodRoutingTests(unittest.TestCase):
    def test_cumulative_import_preserves_weekly_categories_and_organizations(self):
        self.check_routing(date(2026,1,1), True)
    def test_weekly_import_preserves_cumulative_report(self):
        self.check_routing(date(2026,10,15), False)
    def check_routing(self,start,cumulative):
        with tempfile.TemporaryDirectory() as tmp:
            app=Path(tmp); previous={'total':1,'period':'prior','items':[]}
            for name in ['error-categories.json','error-cumulative.json','error-organizations.json']:
                (app/name).write_text(json.dumps(previous))
            (app/'mo-registry.json').write_text('{"organizations":[]}')
            book=Workbook(); sheet=book.active
            sheet.cell(7,2,'1.2.3');sheet.cell(7,3,'Example MO');sheet.cell(7,9,'Error');sheet.cell(7,10,7)
            source=app/'errors.xlsx';book.save(source)
            result=adapt_errors(app,source,date(2026,10,23),start)
            target='error-cumulative.json' if cumulative else 'error-categories.json'
            self.assertEqual(json.loads((app/target).read_text())['total'],7)
            retained='error-categories.json' if cumulative else 'error-cumulative.json'
            self.assertEqual(json.loads((app/retained).read_text()),previous)
            if cumulative:self.assertEqual(json.loads((app/'error-organizations.json').read_text()),previous)
            self.assertIn(target,result.changed_files if hasattr(result,'changed_files') else result.changedFiles)

if __name__=='__main__':unittest.main()
