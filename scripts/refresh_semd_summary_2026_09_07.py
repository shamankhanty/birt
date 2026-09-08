#!/usr/bin/env python3
"""Refresh the regional REMD summary from the approved 07.09.2026 source."""
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = Path("/workspace/scratch/e3dedb97a4e9/vks-0709-extracted")
SOURCE = next(path for path in SOURCE_ROOT.rglob("*.xlsx") if "Медкнижки" in path.name)

values_wb = load_workbook(SOURCE, read_only=True, data_only=True)
formula_wb = load_workbook(SOURCE, read_only=True, data_only=False)
values_ws = values_wb["Отчет РЭМД"]
formula_ws = formula_wb["Отчет РЭМД"]
names = next(formula_ws.iter_rows(min_row=5, max_row=5, values_only=True))
formats = next(formula_ws.iter_rows(min_row=6, max_row=6, values_only=True))
values = next(values_ws.iter_rows(min_row=8, max_row=8, values_only=True))

combined = defaultdict(lambda: {"count": 0, "formats": []})
last_name = ""
for column in range(4, values_ws.max_column + 1):
    if names[column - 1]:
        last_name = str(names[column - 1]).strip()
    if not last_name:
        continue
    value = values[column - 1]
    combined[last_name]["count"] += int(value) if isinstance(value, (int, float)) else 0
    item_format = str(formats[column - 1]).strip() if formats[column - 1] else ""
    if item_format and item_format not in combined[last_name]["formats"]:
        combined[last_name]["formats"].append(item_format)

payload = {
    "period": "01.01.2026–07.09.2026",
    "formed": "07.09.2026 11:08",
    "total": int(values[2]),
    "registeredTypes": sum(item["count"] > 0 for item in combined.values()),
    "items": [
        {
            "name": name,
            "count": item["count"],
            "format": " / ".join(item["formats"]) or "—",
        }
        for name, item in combined.items()
    ],
}
assert sum(item["count"] for item in payload["items"]) == payload["total"]
(ROOT / "app" / "semd-summary.json").write_text(
    json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)
print(json.dumps({key: payload[key] for key in ("period", "formed", "total", "registeredTypes")}, ensure_ascii=False))
