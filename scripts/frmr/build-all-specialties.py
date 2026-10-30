"""Stage all-specialty FRMR/FRMO extracts. Never export full employee identifiers."""
import json,re,hashlib,collections,datetime,sys
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[2];src=Path(sys.argv[1]);cut=datetime.date(2026,10,29);prefix='1.2.643.5.1.13.13.12.2.16.'
clean=lambda v:str(v or '').strip()
def write(name,data): (ROOT/'app'/name).write_text(json.dumps(data,ensure_ascii=False,separators=(',',':'))+'\n')
def source(p):return {'name':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
f=next(src.glob('АНАЛИЗ*.xlsx'));p=next(src.glob('Привязка*.xlsx'));wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
ws=wb['Данные из ФРМР'];headers=next(ws.iter_rows(min_row=2,max_row=2,values_only=True));assert headers[12]=='OID организации' and headers[23]=='Должность по фед справочнику'
summary={}
for r in wb['OID МО. СВОД расхождений'].iter_rows(min_row=2,values_only=True):
 if r[2] and r[13]:
  oid=clean(r[13]);assert oid.startswith(prefix) and oid not in summary
  assert isinstance(r[19],(int,float));summary[oid]={'name':clean(r[2]),'count':int(r[19]),'state':r[3]==1}
mos=sorted(summary);mi={oid:i for i,oid in enumerate(mos)};roles=[];sp=[];ri={};si={};records=[];employment=[];group=collections.defaultdict(lambda:[0,0]);by=collections.defaultdict(collections.Counter);specialties=collections.Counter();quality=collections.Counter();boundary=[];sourceIssueSum=0
for rownum,r in enumerate(ws.iter_rows(min_row=3,values_only=True),3):
 if not r[12] or not r[23]:continue
 oid=clean(r[12]);role=clean(r[23]);spoid=clean(r[16]);name=clean(r[19]) or 'Наименование СП не указано'
 assert oid in mi and r[38]==1
 assert not spoid or spoid.startswith(oid+'.0.'),f'Unsupported SP OID at row {rownum}'
 if role not in ri:ri[role]=len(roles);roles.append(role)
 key=(oid,spoid)
 if key not in si:si[key]=len(sp);sp.append([mi[oid],spoid.removeprefix(oid+'.0.') if spoid else '',name])
 m,a,b=mi[oid],ri[role],si[key];stage1=clean(r[34]).startswith('1 этап.');stage2=clean(r[33]).startswith('2 этап.');assert stage1 or stage2
 issue=0 if stage1 else 1;digits=re.sub(r'\D','',clean(r[5]));last4=digits[-4:] if len(digits)==11 else None;flags=0
 if last4 is None:flags|=1;quality['missingIdentifier']+=1
 if isinstance(r[27],datetime.datetime) and r[27].date()<=cut:flags|=2;quality['employmentEnded']+=1
 if clean(r[32])=='НЕТ СВЯЗКИ С ФРМО':flags|=8;quality['missingOrganizationLink']+=1
 records.append([m,a,b,last4,issue,rownum,flags]);employment.append(tuple(clean(r[k]) for k in [5,12,16,20,22,23,25,26,27]));group[(m,a,b)][issue]+=1
 by[oid]['current']+=1;by[oid]['stage1' if issue==0 else 'stage2']+=1;by[oid]['role:'+role]+=1;specialties[role]+=1;sourceIssueSum+=r[37]
 boundary.append(rownum)
counts=collections.Counter(employment)
for rec,key in zip(records,employment):
 if counts[key]>1:rec[6]|=4;quality['repeatedRecordRows']+=1
quality['repeatedRecordExtras']=sum(n-1 for n in counts.values() if n>1)
assert len(records)==sum(r['count'] for r in summary.values())==8264
assert sourceIssueSum==13878
for oid,v in summary.items():assert v['count']==by[oid]['current'],oid
rows=[]
for oid in mos:
 v=summary[oid];b=by[oid];specs=sorted([[k[5:],n] for k,n in b.items() if k.startswith('role:')],key=lambda r:(-r[1],r[0]))
 rows.append([oid.removeprefix(prefix),v['name'],None,None,v['count'],b['stage1'],b['stage2'],None,specs])
summaryOut={'schemaVersion':2,'source':f.name,'sourceCut':'29.10.2026','baselineCut':None,'scope':'all_specialties','unit':'должностная запись, не уникальный сотрудник','trend':[['29.10.2026',len(records)]],'topSpecialties':sorted(specialties.items(),key=lambda r:(-r[1],r[0])),'rows':rows}
mosuffix=[o.removeprefix(prefix) for o in mos]
out={'summary':summaryOut,'departments':{'cut':'29.10.2026','unit':'должностная запись','mos':mosuffix,'roles':roles,'sp':sp,'rows':[[*key,*value] for key,value in sorted(group.items())]},'records':{'cut':'29.10.2026','issues':['Должность не найдена на уровне МО в ГИС','Должность есть в МО, но не найдена в указанном СП'],'qualityBits':{'1':'Отсутствует корректный идентификатор в исходнике','2':'Дата окончания работы наступила к дате среза','4':'Совпадают проверенные реквизиты кадровых записей; требуется проверка повтора','8':'В служебном поле источника нет связи МО с ФРМО'},'rows':records},'metadata':{'source':source(f),'sourceRowRange':[boundary[0],boundary[-1]],'parsedRows':len(records),'sourceIssueCount':sourceIssueSum,'quality':quality,'scopeNote':'Все должности из представленного списка расхождений. Общая численность всех кадровых записей в выгрузке не представлена; доля ошибок не рассчитывается. История 22.10 имеет другой охват.','stateOrganizations':sum(v['state'] for v in summary.values()),'stateProblemRecords':sum(v['count'] for v in summary.values() if v['state'])}}
write('frmr-all-specialties.json',out)
# Source FRMO summary counts department rows, not validated deviations.
wb=openpyxl.load_workbook(p,read_only=True,data_only=True);s=wb['Отделения и СП (филиалы)'];h=next(s.iter_rows(min_row=2,max_row=2,values_only=True));assert h[0]=='Код МО' and 'ОИД' in h[4]
sourceSummary={clean(r[0]):[clean(r[1]),r[2],r[3]] for r in wb['СВОД на актуальную дату'].iter_rows(min_row=4,values_only=True) if r[0] and r[1]}
org={};detail=[];dates=[];flagcounts=collections.Counter();deptrows=collections.Counter();parents=collections.defaultdict(set)
for rownum,r in enumerate(s.iter_rows(min_row=4,values_only=True),4):
 if not r[0] or not r[1] or not r[6]:continue
 code=clean(r[0]);org[code]=clean(r[1]);b,d=clean(r[4]),clean(r[8]);flags=0
 for value in [b,d]:
  assert not value or re.fullmatch(r'\d+(\.\d+)+',value)
  if value:parents[code].add(value.split('.0.')[0])
 if not b:flags|=1
 if not d:flags|=2
 if b and d and b!=d:flags|=4
 if clean(r[3])!=clean(r[7]):flags|=8
 if clean(r[5])=='Нет' and clean(r[9])=='1':flags|=16
 date=r[10].date().isoformat() if isinstance(r[10],datetime.datetime) else clean(r[10]) or None
 detail.append([code,clean(r[2]),clean(r[6]),b or None,d or None,flags,rownum,date]);deptrows[code]+=1
 for bit in [1,2,4,8,16]:flagcounts[str(bit)]+=int(bool(flags&bit))
for code,(name,current,previous) in sourceSummary.items():assert current==deptrows[code],(code,current,deptrows[code])
missing=sorted(set(org)-set(sourceSummary));assert sum(deptrows[c] for c in missing)==34
flags={'1':'Нет OID СП у филиала','2':'Нет OID СП у отделения','4':'Разные OID СП у филиала и отделения','8':'Разные привязанные наименования ФРМО','16':'Активное отделение в неактивном СП по выгрузке'}
result={'cut':'30.10.2026','source':source(p),'flags':flags,'rows':detail,'organizations':[{'code':c,'name':org[c],'rows':deptrows[c],'issues':sum(bool(r[5]) for r in detail if r[0]==c),'sourceSummary':sourceSummary.get(c,[None,None,None])[1],'summaryMissing':c not in sourceSummary,'parentOids':sorted(parents[c])} for c in sorted(org)],'totals':{'rows':len(detail),'issues':sum(bool(r[5]) for r in detail),'consistent':sum(not r[5] for r in detail),'sourceSummaryRows':sum(v[1] for v in sourceSummary.values()),'missingSummaryCodes':len(missing),'missingSummaryRows':sum(deptrows[c] for c in missing),'flagCounts':flagcounts},'method':'Одна строка — отделение ГИС. Строка с хотя бы одним флагом учитывается один раз. Различие наименований требует проверки справочников и не доказывает ошибку OID. Категории пересекаются. Сводка исходника считает строки, поэтому её значение не используется как число ошибок.','comparison':'Предыдущая детализация с теми же правилами не представлена; динамика ошибок не рассчитывается.'}
assert result['totals']['rows']==3385 and result['totals']['issues']==2043
write('frmo-linkage-data.json',result)
print(json.dumps({'frmr':{'positions':len(records),'organizations':len(mos),'roles':len(roles),'groups':len(group),'quality':quality},'frmo':result['totals']},ensure_ascii=False))
