#!/usr/bin/env python3
import json
import re
import subprocess
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SHORT = Path('/workspace/scratch/1f5740e14d75/upload/Случаи краткого ввода 01.01.-28.08. (Выгрузка из ГИС ЭЗ РТ 28.08.26).xlsx')
EPL = Path('/workspace/scratch/1f5740e14d75/upload/Отчёт_по_использованию_системы_24_08_30_08.xlsx')
EPL_PREVIOUS = Path('/workspace/scratch/1f5740e14d75/library-downloads/Отчёт_по_использованию_системы_17_08_2026_по_23_08_2026 (2).xlsx')

EPL_GROUP_NAMES = {
    110: 'ССМП Казани (подстанции №1–9 и общая строка)',
    118: 'База медицинского обеспечения (филиалы и головная организация)',
    119: 'РКПБ им. В.М. Бехтерева (филиалы и головная организация)',
    120: 'Республиканский клинический КВД (филиалы и головная организация)',
    121: 'Республиканский клинический наркологический диспансер (филиалы и головная организация)',
    122: 'РКОД им. М.З. Сигала (филиалы и головная организация)',
    123: 'Противотуберкулёзная служба РТ (7 объектов)',
    124: 'Республиканский центр крови (филиалы и головная организация)',
    125: 'Центр СПИД РТ (филиалы и головная организация)',
    126: 'Республиканское бюро СМЭ (отделения и головная организация)',
}

def n(v):
    return int(v or 0)

def norm(v):
    return re.sub(r'\s+', ' ', str(v or '').strip()).casefold()

def update_short_input():
    path = ROOT / 'app' / 'operational-mo.json'
    data = json.loads(path.read_text(encoding='utf-8'))
    baseline = json.loads(subprocess.check_output(['git','show','HEAD:app/operational-mo.json'], cwd=ROOT, text=True))
    old = baseline['shortInput']['rows']
    old_by_key = {(norm(r.get('name')), r.get('oid') or ''): r for r in old}
    ws = load_workbook(SHORT, read_only=True, data_only=True)['Краткий ввод']
    raw = []
    for values in ws.iter_rows(min_row=7, values_only=True):
        if not values[0] or str(values[0]).strip().lower().startswith('итого'):
            continue
        name, oid = str(values[0]).strip(), str(values[1] or '').strip()
        amb, round_hosp, day_hosp, preventive = n(values[3]), n(values[4]), n(values[5]), n(values[6])
        raw.append((name, oid, amb + preventive, round_hosp + day_hosp))

    def rows_for(kind):
        out=[]
        for name, oid, amb, hosp in raw:
            fact = amb + hosp if kind == 'shortInput' else amb if kind == 'shortInputAmb' else hosp
            old_row = old_by_key.get((norm(name), oid))
            previous = old_row.get('fact', 0) if old_row else None
            row={'name':name,'oid':oid,'fact':fact,'count':fact,'previous':previous,'trend':fact-previous if previous is not None else None}
            if previous is not None and fact < previous:
                row['sourceWarning']='Накопительное значение уменьшилось; корректировка источника требует уточнения.'
            out.append(row)
        return out

    for key, title in [('shortInput','Количество случаев краткого ввода — всего'),('shortInputAmb','Количество случаев краткого ввода — амбулаторно'),('shortInputHosp','Количество случаев краткого ввода — стационар')]:
        data[key].update({'name':title,'date':'28.08.2026','period':'01.01–28.08.2026','rows':rows_for(key),'mode':'count','direction':'lower','note':'Справочный накопительный показатель без норматива: не влияет на рейтинг и приоритет заслушивания. Уменьшение накопительного значения отмечается как риск качества источника.'})
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    return {k:sum(r['fact'] for r in data[k]['rows']) for k in ('shortInput','shortInputAmb','shortInputHosp')}

def extract_epl(source, period):
    ws = load_workbook(source, read_only=True, data_only=True)['Лист1']
    groups=[]
    current=None
    fields=('vehicles','ambulanceVehicles','otherVehicles','moved','drivers','mechanics','medics','waybills','ambulanceWaybills','otherWaybills','driversWithWaybills')
    for v in ws.iter_rows(min_row=14, values_only=True):
        if not v[1]:
            continue
        if str(v[1]).strip().casefold() == 'итого':
            continue
        if isinstance(v[0], (int,float)):
            source_number=n(v[0])
            current={'sourceNumber':source_number,'name':EPL_GROUP_NAMES.get(source_number, str(v[1]).strip()), **{f:0 for f in fields}, 'components':[]}
            groups.append(current)
        elif current is None:
            continue
        vals=[n(v[2]),n(v[3]),n(v[4]),n(v[5]),n(v[6]),n(v[7]),n(v[8]),n(v[9]),n(v[10]),n(v[11]),n(v[12])]
        for f,x in zip(fields,vals): current[f]+=x
        component={'name':str(v[1]).strip()}
        for f,x in zip(fields,vals): component[f]=x
        component['movementShare']=component['moved']/component['vehicles']*100 if component['vehicles'] else None
        current['components'].append(component)
    for r in groups:
        r['movementShare']=r['moved']/r['vehicles']*100 if r['vehicles'] else None
        if len(r['components']) == 1:
            r['components'] = []
    vehicles=sum(r['vehicles'] for r in groups); moved=sum(r['moved'] for r in groups)
    detail={'organizations':len(groups),'vehicles':vehicles,'vehiclesWithMovement':moved,'movementShare':moved/vehicles*100,'waybills':sum(r['waybills'] for r in groups),'driversWithWaybills':sum(r['driversWithWaybills'] for r in groups),'organizationsWithMovement':sum(r['vehicles']>0 and r['moved']>0 for r in groups),'zeroMovementOrganizations':sum(r['vehicles']>0 and r['moved']==0 for r in groups),'zeroVehicleOrganizations':sum(r['vehicles']==0 for r in groups)}
    summary={'organizations':n(ws['C6'].value),'vehicles':n(ws['C7'].value),'vehiclesWithWaybills':n(ws['E7'].value),'vehiclesWithMovement':n(ws['H7'].value)}
    if vehicles != summary['vehicles'] or moved != summary['vehiclesWithMovement']:
        raise ValueError(f'ЭПЛ {period}: детализация не совпадает с итогом источника: ТС {vehicles}/{summary["vehicles"]}, с движением {moved}/{summary["vehiclesWithMovement"]}')
    return {'source':source.name,'period':period,'sourceHeading':str(ws['B2'].value or ''),'systemSummary':summary,'detail':detail,'rows':groups}

def update_epl():
    path = ROOT / 'app' / 'electronic-waybill-weekly.json'
    previous=extract_epl(EPL_PREVIOUS, '17.08–23.08.2026')
    current_data=extract_epl(EPL, '24.08–30.08.2026')
    out={'previous':previous,'current':current_data,'comparisonRule':'Недельная динамика рассчитана по одинаковым 126 группам источника. Дочерние строки без номера включены в соответствующую группу; сравнение выполняется по номеру группы, а подразделения доступны в расшифровке.','periodCorrection':'Обе недели повторно разобраны по единому правилу: 126 управленческих групп и 171 строка подразделений.'}
    path.write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
    return current_data['detail'], current_data['systemSummary']

if __name__ == '__main__':
    print(json.dumps({'shortInput':update_short_input(),'epl':update_epl()},ensure_ascii=False,indent=2))
