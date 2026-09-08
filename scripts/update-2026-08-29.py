#!/usr/bin/env python3
"""Load the 29.08 hospital/FAP reports and the missed 28.08 ASU SMP report."""
from __future__ import annotations

import json
import re
import subprocess
from collections import defaultdict
from pathlib import Path

import openpyxl
import xlrd

ROOT = Path("/workspace/sites/rt-health-digital-dashboard")
APP = ROOT / "app"
UPLOAD = Path("/workspace/scratch/1f5740e14d75/upload")
AUDIT = Path("/workspace/scratch/1f5740e14d75/audit_28_08_pjuDug")
HOSPITAL = UPLOAD / "Отчет_по_госпитализациям__-_01.01.2026-29.08.2026.xlsx"
FAP = UPLOAD / "Отчет_по_ФАП_и_ФП_(кол-во_зарегистрированных_СЭМД)_-_01.01.2026-29.08.2026.xlsx"
SMP = AUDIT / "Отправка сведений в федеральные сервисы 01.01.-28.08. (Выгрузка из АСУ ССМП 28.08.26).xls"


def load(name: str):
    return json.loads((APP / name).read_text(encoding="utf-8"))


def save(name: str, data):
    (APP / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def number(value) -> int:
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0


def norm(value: str) -> str:
    value = str(value or "").lower().replace("ё", "е")
    value = re.sub(r'["«»]', "", value)
    value = re.sub(r"^(?:гауз|гбуз|гбу|фгбу|фгаоу\s*во|ао|ооо)(?:\s+рт)?\s+", "", value)
    value = re.sub(r"\s*\([^)]*\)\s*$", "", value)
    return re.sub(r"[^а-яa-z0-9№]+", " ", value).strip()


mo = load("mo-data.json")
details = load("mo-details.json")
detail_oids = load("mo-detail-oids.json")
operational = load("operational-mo.json")
units = load("unit-data.json")
organization_status = load("organization-status.json")
base_mo = json.loads(subprocess.check_output(["git", "show", "HEAD:app/mo-data.json"], cwd=ROOT, text=True))
base_operational = json.loads(subprocess.check_output(["git", "show", "HEAD:app/operational-mo.json"], cwd=ROOT, text=True))


def previous_map(metric: str):
    result = {}
    source = base_operational if metric in base_operational else base_mo
    for row in source.get(metric, {}).get("rows", []):
        if row.get("oid"):
            result[("oid", str(row["oid"]))] = row.get("fact")
        result[("name", norm(row.get("name")))] = row.get("fact")
    return result


# Hospital: denominator is sheet 3; numerator is the sum of both discharge documents.
old_hospital = previous_map("hospital")
wb = openpyxl.load_workbook(HOSPITAL, read_only=True, data_only=True)
hospital_rows = [row for row in wb["Лист3"].iter_rows(min_row=6, values_only=True) if row[1]]
rows, by_name, by_oid = [], {}, {}
for source in hospital_rows:
    name, oid = str(source[1]), str(source[2])
    volume = number(source[3])
    registered = number(source[4]) + number(source[5])
    fact = registered / volume * 100 if volume else 0
    previous = old_hospital.get(("oid", oid), old_hospital.get(("name", norm(name))))
    row = {"name": name, "oid": oid, "fact": fact, "count": registered,
           "previous": previous, "trend": None if previous is None else fact - previous}
    if fact > 100:
        row["sourceWarning"] = "Значение выше 100%: требуется проверка повторных/исправленных СЭМД."
    rows.append(row)
    value = {"volume": volume, "registered": registered}
    by_name[norm(name)] = value
    by_oid[oid] = value
mo["hospital"] = {"name": mo["hospital"]["name"], "plan": 95, "unit": "%",
                  "date": "29.08.2026", "period": "01.01–29.08.2026",
                  "note": "Знаменатель — лист 3. Фактические значения свыше 100% показаны без искажения, рейтинг ограничивается 100 баллами.",
                  "rows": rows}
details["hospital"], detail_oids["hospital"] = by_name, by_oid
operational["hospitalCount"] = {"name": "Количество госпитализаций", "plan": None, "unit": "",
    "date": "29.08.2026", "period": "01.01–29.08.2026", "mode": "count",
    "rows": [{"name": str(r[1]), "oid": str(r[2]), "fact": number(r[3]), "count": number(r[3]),
              "previous": None, "trend": None} for r in hospital_rows]}


# FAP/FP: sheet 1 is the full list (including zero units); aggregate by medical organization.
old_fap = previous_map("fapSemdCount")
wb = openpyxl.load_workbook(FAP, read_only=True, data_only=True)
grouped = defaultdict(lambda: {"count": 0, "units": 0, "zero": 0})
for source in wb["Лист1"].iter_rows(min_row=5, values_only=True):
    if not source[0]:
        continue
    group = grouped[str(source[0])]
    value = number(source[3])
    group["count"] += value
    group["units"] += 1
    group["zero"] += int(value == 0)
fap_rows = []
for name, value in sorted(grouped.items(), key=lambda item: norm(item[0])):
    previous = old_fap.get(("name", norm(name)))
    row = {"name": name, "fact": value["count"], "count": value["count"], "previous": previous,
           "trend": None if previous is None else value["count"] - previous,
           "units": value["units"], "zeroUnits": value["zero"]}
    if value["zero"]:
        row["sourceWarning"] = f'{value["zero"]} ФАП/ФП с нулевым результатом из {value["units"]}'
    fap_rows.append(row)
operational["fapSemdCount"] = {"name": "Количество зарегистрированных СЭМД по ФАП и ФП", "plan": None,
    "unit": "", "date": "29.08.2026", "period": "01.01–29.08.2026", "mode": "count",
    "note": "Лист 1 содержит полный перечень ФАП/ФП; нулевые подразделения сохранены и показаны отдельно.",
    "rows": fap_rows}


# SMP: keep only organization total rows. Identical station/detail duplicates are removed.
old_smp = previous_map("smp")
book = xlrd.open_workbook(str(SMP), formatting_info=True)
sheet = book.sheet_by_index(0)
smp_rows = []
for index in range(8, sheet.nrows):
    source = [sheet.cell_value(index, col) for col in range(sheet.ncols)]
    name = str(source[0]).strip()
    cell = sheet.cell(index, 0)
    is_organization_total = bool(book.font_list[book.xf_list[cell.xf_index].font_index].bold)
    if not name or not is_organization_total:
        continue
    volume, registered = number(source[1]), number(source[10])
    fact = registered / volume * 100 if volume else 0
    previous = old_smp.get(("name", norm(name)))
    smp_rows.append({"name": name, "fact": fact, "count": registered, "volume": volume,
                     "registered": registered, "previous": previous,
                     "trend": None if previous is None else fact - previous})
mo["smp"] = {"name": mo["smp"]["name"], "plan": 95, "unit": "%", "date": "28.08.2026",
             "period": "01.01–28.08.2026",
             "note": "Числитель — статус «Принято» АСУ СМП, принимаемый как регистрация в РЭМД; знаменатель — количество карт вызова АСУ СМП за тот же период.",
             "rows": smp_rows}
details["smp"] = {norm(r["name"]): {"volume": r["volume"], "registered": r["registered"]} for r in smp_rows}


# Federal building-level reports: refresh both the unit register and the MO aggregation.
def refresh_building_metric(metric: str, path: Path, sheet_name: str, start_row: int, title: str, entity: str):
    old = previous_map(metric)
    old_counts = {(str(r.get("moOid", "")), str(r.get("buildingIds", r.get("unitOid", "")))): r.get("count", 0)
                  for r in units.get(metric, {}).get("rows", [])}
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    source_rows = []
    grouped = defaultdict(lambda: {"name": "", "volume": 0, "registered": 0})
    for source in book[sheet_name].iter_rows(min_row=start_row, values_only=True):
        if source[1] != "Республика Татарстан" or not source[2]:
            continue
        name, oid, building_id = str(source[2]), str(source[3]), str(source[4])
        registered = str(source[6]).strip().lower() == "да"
        source_rows.append({"mo": name, "moOid": oid, "unit": str(source[5] or "Наименование объекта отсутствует"),
                            "unitOid": building_id, "buildingIds": building_id, "registered": registered,
                            "count": old_counts.get((oid, building_id), 0)})
        group = grouped[oid]
        group["name"] = name
        group["volume"] += 1
        group["registered"] += int(registered)
    units[metric] = {"name": title, "entity": entity, "plan": len(source_rows),
                     "fact": sum(r["registered"] for r in source_rows), "date": "28.08.2026", "rows": source_rows}
    metric_rows, metric_details, metric_oids = [], {}, {}
    for oid, group in grouped.items():
        fact = group["registered"] / group["volume"] * 100 if group["volume"] else 0
        previous = old.get(("oid", oid), old.get(("name", norm(group["name"]))))
        metric_rows.append({"name": group["name"], "oid": oid, "fact": fact, "count": group["registered"],
                            "previous": previous, "trend": None if previous is None else fact - previous})
        value = {"volume": group["volume"], "registered": group["registered"]}
        metric_details[norm(group["name"])] = value
        metric_oids[oid] = value
    mo[metric] = {"name": mo[metric]["name"], "plan": 100, "unit": "%", "date": "28.08.2026",
                  "period": "01.01–27.08.2026", "rows": metric_rows}
    details[metric], detail_oids[metric] = metric_details, metric_oids


refresh_building_metric(
    "tvspDiagnostic",
    next(AUDIT.glob("Доля_ТВСП*.xlsx")),
    "Факт передачи", 8,
    "ТВСП, передающие протоколы диагностических исследований", "объект контроля",
)
refresh_building_metric(
    "smpFederal",
    next(AUDIT.glob("Отчет_СМП_ТВСП*.xlsx")),
    "Факт передачи", 7,
    "Станции и подстанции СМП, передающие карты вызова", "станция / подстанция СМП",
)


# TMK REMD: the detailed report contains 39 transmitting organizations; retain the approved list of 40.
tmk_path = next(AUDIT.glob("Отчет_ТМК_РЭМД*.xlsx"))
book = openpyxl.load_workbook(tmk_path, read_only=True, data_only=True)
counts = {str(r[3]): number(r[4]) for r in book["Детализированный отчет"].iter_rows(min_row=7, values_only=True)
          if r[1] == "Республика Татарстан" and r[3]}
tmk_rows = []
for old in organization_status["tmkRemd"]["rows"]:
    count = counts.get(str(old.get("oid")), 0)
    fact = 100 if count > 0 else 0
    previous = old.get("fact")
    tmk_rows.append({**old, "count": count, "fact": fact, "previous": previous,
                     "trend": None if previous is None else fact - previous})
organization_status["tmkRemd"].update({"date": "28.08.2026", "period": "01.01–27.08.2026",
    "note": "Плановый перечень — 40 МО; 39 МО передают протоколы ТМК в РЭМД.", "rows": tmk_rows})
organization_status["elmk"]["period"] = "01.01–28.08.2026"

save("mo-data.json", mo)
save("mo-details.json", details)
save("mo-detail-oids.json", detail_oids)
save("operational-mo.json", operational)
save("unit-data.json", units)
save("organization-status.json", organization_status)

summary = {
    "hospital": {"rows": len(rows), "volume": sum(v["volume"] for v in by_oid.values()),
                 "registered": sum(v["registered"] for v in by_oid.values())},
    "fap": {"organizations": len(fap_rows), "units": sum(v["units"] for v in grouped.values()),
            "zero_units": sum(v["zero"] for v in grouped.values()), "documents": sum(v["count"] for v in grouped.values())},
    "smp": {"organizations": len(smp_rows), "volume": sum(r["volume"] for r in smp_rows),
            "registered": sum(r["registered"] for r in smp_rows)},
}
print(json.dumps(summary, ensure_ascii=False, indent=2))
