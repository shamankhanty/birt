"""Approved correction: use the same-date general REMD export for 122/228."""
import json
from datetime import date
from pathlib import Path
import openpyxl
from pipeline.adapters import adapt_preventive, parse_remd, parse_foms

root=Path(__file__).resolve().parents[1]
sources=root.parent/'intake-20261014/input'
remd=next(sources.glob('1.Отчет*'))
foms=next(sources.glob('Количество_СЭМД_Результаты*'))
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def write(p,d):p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
for p in (remd,foms):
    w=openpyxl.load_workbook(p,read_only=True,data_only=True)
    heading=' '.join(str(v)for r in list(w.worksheets[0].values)[:2]for v in r if v)
    assert '01.01.2026' in heading and '13.10.2026' in heading
rr,ff=parse_remd(remd),parse_foms(foms)
mo_source=openpyxl.load_workbook(remd,read_only=True,data_only=True)['Отчет РЭМД по МО']
for row in mo_source.iter_rows(min_row=8,values_only=True):
    if row[2] in ff:
        assert isinstance(row[80],(int,float)) and isinstance(row[107],(int,float))
missing=set(ff)-set(rr)
assert missing=={'1.2.643.5.1.13.13.12.2.18.1384'}
assert ff[next(iter(missing))]['den']==1
assert len(set(ff)&set(rr))==95
assert sum(max(rr[k]['s122'],rr[k]['s228'])for k in ff if k in rr)==1915517
assert sum(v['den']for k,v in ff.items()if k in rr)==2281382
result=adapt_preventive(root/'app',remd,foms,date(2026,10,13),root/'app/mo-registry.json',missing)
audit=read(root/'app/preventive-semd-audit.json')
assert audit['summary']['numerator']==1915517 and audit['summary']['denominator']==2281382
methodology=read(root/'app/methodology-audit.json')
item=next(r for r in methodology['items'] if r['source'].startswith('Профилактика'))
item.update(status='Исправлена',issue='Найдена сопоставимая пара источников за 01.01–13.10.2026',impact='Ранее общий отчет РЭМД не выбирался как источник СЭМД 122/228, и показатель оставался на 06.10.',fix='MAX(122;228) рассчитан отдельно по каждой из 95 МО; итог 1 915 517 / 2 281 382 = 83,96%. Внешняя строка из Ижевска (знаменатель 1) без соответствия РЭМД сохранена в аудите и не заменена нулем.')
if not any(r['source']=='Профилактика, 01.01–13.10.2026' for r in methodology['items']):
    methodology['items'].append(dict(severity='high',source='Профилактика, 01.01–13.10.2026',status='Требует уточнения',issue=f"У {audit['summary']['over100']} МО выбранное количество СЭМД превышает обращения ФОМС",impact='Доля источника выше 100%; возможны различия состава обращений и зарегистрированных документов.',fix='Фактические компоненты сохранены без ограничения и без подстановки нулей. Требуется сверка уникальных обращений и повторных версий СЭМД по МО.'))
write(root/'app/methodology-audit.json',methodology)
for p in (root/'app').glob('*.json'):
    if p.name in result.changedFiles:write(p,read(p))
write(root/'docs/preventive-20261013.json',{'sources':[remd.name,foms.name],'period':'01.01–13.10.2026','summary':audit['summary'],'excluded':audit['excluded'],'facts':result.facts,'previousPeriod':'01.01–06.10.2026','formulaUnchanged':True})
print(json.dumps(audit['summary'],ensure_ascii=False))
