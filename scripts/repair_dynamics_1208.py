#!/usr/bin/env python3
"""Repair 12.08 dynamics from the immutable pre-update snapshot.

The script is intentionally deterministic and leaves non-comparable/new-baseline
datasets without a fabricated previous value.
"""
import json, re, subprocess, glob
from collections import defaultdict
from pathlib import Path
import xlrd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
UP=Path('/workspace/scratch/8e713f374d53/upload')

def load(name): return json.loads((ROOT/'app'/name).read_text(encoding='utf-8'))
def save(name,obj): (ROOT/'app'/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
def old(name): return json.loads(subprocess.check_output(['git','show',f'HEAD^:app/{name}']))
def text(v): return '' if v is None else str(v).strip()
def norm(v):
    v=text(v).lower().replace('ё','е')
    v=re.sub(r'^(?:филиал\s+)?(?:гауз|гбуз|гбу|фгбу|фгаоу\s*во|ао|ооо|чуз)(?:\s+рт)?\s*','',v)
    v=re.sub(r'["«»]','',v);v=re.sub(r'\s*\([^)]*\)\s*$','',v)
    v=v.replace('городская клиническая больница','гкб').replace('городская больница','гб').replace('городская поликлиника','гп').replace('центральная районная больница','црб')
    v=re.sub(r'\s*№\s*','№',v)
    return re.sub(r'[^а-яa-z0-9№]+',' ',v).strip()

registry=load('mo-registry.json')
alias={}
for org in registry['organizations']:
    for a in org.get('aliases',[]):
        k=norm(a); alias.setdefault(k,set()).add(org['oid'])
def identity(row):
    if row.get('oid'): return 'oid:'+text(row['oid'])
    ids=alias.get(norm(row.get('name')),set())
    if len(ids)==1:return 'oid:'+next(iter(ids))
    return 'name:'+norm(row.get('name'))
def old_index(rows):
    out={}
    for r in rows:
        for k in (identity(r),'name:'+norm(r.get('name'))):
            if k not in out:out[k]=r
    return out
def restore(ds,prev_ds,comparable=True):
    idx=old_index(prev_ds.get('rows',[]))
    for r in ds['rows']:
        p=(idx.get(identity(r)) or idx.get('name:'+norm(r.get('name')))) if comparable else None
        r['previous']=p.get('fact') if p else None
        r['trend']=r['fact']-r['previous'] if r['previous'] is not None else None

mo=load('mo-data.json'); old_mo=old('mo-data.json')
op=load('operational-mo.json'); old_op=old('operational-mo.json')
details=load('mo-details.json')
units=load('unit-data.json'); old_units=old('unit-data.json')
status=load('organization-status.json')

# Comparable MO indicators. Hospital stays incomparable after the denominator change.
for key in ['egpu','egpu2days','birth','death','semd228','ambulatoryCase']:
    restore(mo[key],old_mo[key])
# Three birth facilities changed from established abbreviations to full names.
birth_aliases={
 'ГАУЗ "Клиника медицинского университета" г.Казани':'КМУ г.Казани',
 'ГАУЗ "Городская детская больница №1" г.Казани':'ГДБ №1 г.Казани',
 'ГАУЗ "Детская городская поликлиника №4" г. Казани':'ДГП №4 г. Казани',
}
birth_old={norm(r['name']):r for r in old_mo['birth']['rows']}
for r in mo['birth']['rows']:
    alias_name=birth_aliases.get(r['name'])
    if alias_name and norm(alias_name) in birth_old:
        r['previous']=birth_old[norm(alias_name)]['fact'];r['trend']=r['fact']-r['previous']
for r in mo['hospital']['rows']:
    r['previous']=None;r['trend']=None

# Remove the explanatory legend accidentally parsed as a birth organization.
mo['birth']['rows']=[r for r in mo['birth']['rows'] if not r['name'].startswith('Где в столбце')]
details['birth']={k:v for k,v in details['birth'].items() if not k.startswith('где в столбце')}

# Reaggregate current/previous unit-level datasets to MO level and restore true dynamics.
unit_titles={
 'tvspStationary':'Доля ТВСП, передающих выписные эпикризы',
 'tvspAmbulatory':'Доля амбулаторных ТВСП, передающих эпикриз/талон и/или протокол консультации',
 'tvspLaboratory':'Доля КДЛ, передающих протоколы лабораторных исследований',
 'tvspDiagnostic':'Доля ТВСП, передающих протоколы диагностических исследований',
 'smpFederal':'Доля станций и подстанций СМП, передающих карты вызова',
}
for key,title in unit_titles.items():
    def aggregate(rows):
        out=defaultdict(lambda:{'planned':0,'registered':0,'name':''})
        for r in rows:
            oid=text(r.get('moOid')); ident='oid:'+oid if oid else 'name:'+norm(r.get('mo'))
            out[ident]['planned']+=1;out[ident]['registered']+=int(bool(r.get('registered')));out[ident]['name']=text(r.get('mo'));out[ident]['oid']=oid
        return out
    cur=aggregate(units[key]['rows']); prv=aggregate(old_units[key]['rows']); rows=[];det={}
    registry_by_oid={o['oid']:o for o in registry['organizations']}
    for ident,v in cur.items():
        if v.get('oid') in registry_by_oid:v['name']=registry_by_oid[v['oid']]['shortName']
        fact=v['registered']/v['planned']*100 if v['planned'] else 0
        pv=prv.get(ident); previous=pv['registered']/pv['planned']*100 if pv and pv['planned'] else None
        rows.append({'name':v['name'],'oid':v.get('oid',''),'fact':fact,'count':v['registered'],'previous':previous,'trend':fact-previous if previous is not None else None})
        det[norm(v['name'])]={'volume':v['planned'],'registered':v['registered']}
    mo[key]={'name':title,'plan':100,'unit':'%','date':'12.08.2026','period':'01.01–12.08.2026','rows':rows,
             'note':'Доля по МО рассчитана из актуального перечня объектов контроля; детализация ниже приведена по ТВСП/КДЛ/станциям.'}
    details[key]=det

# Operational counts: cumulative growth is shown against the prior snapshot.
for key in ['tmkMaxCount','elnMaxCount']:
    # MAX files have no OID; keep numbered Kazan and Naberezhnye Chelny
    # facilities separate by their exact normalized source names.
    idx={norm(r['name']):r for r in old_op[key]['rows']}
    for r in op[key]['rows']:
        p=idx.get(norm(r['name']));r['previous']=p.get('fact') if p else None;r['trend']=r['fact']-r['previous'] if r['previous'] is not None else None
old_short_file=next(p for p in sorted(glob.glob(str(UP/'*краткого ввода*05.08*.xlsx'))) if Path(p).name.startswith('Случаи краткого ввода'))
old_short_ws=load_workbook(old_short_file,read_only=True,data_only=True).active
old_short={}
for row in old_short_ws.iter_rows(min_row=7,values_only=True):
    oid=text(row[1]);name=text(row[0]);
    old_short[(oid,norm(name))]=sum((row[i] or 0) for i in (3,4,5,6))
for r in op['shortInput']['rows']:
    p=old_short.get((text(r.get('oid')),norm(r.get('name'))));r['previous']=p;r['trend']=r['fact']-p if p is not None else None
for key in ['shortInputAmb','shortInputHosp','fapSemdCount']:
    for r in op[key]['rows']:r['previous']=None;r['trend']=None

# Rebuild ambulance-card MO rows from parent totals only; exact duplicates/substations are excluded.
def find(pattern):
    fs=glob.glob(str(UP/pattern));
    if not fs:raise FileNotFoundError(pattern)
    return sorted(fs,key=lambda p:len(Path(p).name))[0]
def smp_parent_rows(path):
    sh=xlrd.open_workbook(path).sheet_by_index(0); chosen={}
    for i in range(8,sh.nrows):
        row=sh.row_values(i); name=text(row[0]);cards=row[1] if len(row)>10 else 0;accepted=row[10] if len(row)>10 else 0
        parent=bool(re.match(r'^(?:ГАУЗ|ГБУЗ|ГБУ|ФГБУ|АО|ООО|ЧУЗ)\b',name,re.I)) or name in {'Лаишевский','ТЦМК'}
        if not parent or not isinstance(cards,(int,float)) or cards<=0:continue
        k=norm(name)
        if k not in chosen or cards>chosen[k][1]:chosen[k]=(name,cards,accepted)
    return list(chosen.values())
cur_smp=smp_parent_rows(find('*АСУ ССМП 12.08.26*.xls'))
old_smp=smp_parent_rows(find('Отправка сведений в федеральные сервисы 01.01.-05.08*.xls'))
old_by={identity({'name':n}):(n,c,a) for n,c,a in old_smp}
srows=[];sdet={}
for name,cards,accepted in cur_smp:
    pv=old_by.get(identity({'name':name}));fact=accepted/cards*100 if cards else 0
    previous=pv[2]/pv[1]*100 if pv and pv[1] else None
    srows.append({'name':name,'fact':fact,'count':int(accepted),'previous':previous,'trend':fact-previous if previous is not None else None})
    sdet[norm(name)]={'volume':int(cards),'registered':int(accepted)}
if (sum(x[1] for x in cur_smp),sum(x[2] for x in cur_smp))!=(584073,484458):raise RuntimeError('SMP totals do not reconcile')
mo['smp']={'name':'Доля принятых в РЭМД карт вызова СМП','plan':95,'unit':'%','date':'12.08.2026','period':'01.01–12.08.2026','rows':srows,
           'note':'По каждой МО использована родительская итоговая строка источника; повторные строки и строки подстанций не суммируются повторно.'}
details['smp']=sdet

# Presence indicators: unchanged presence means zero change, never -100.
for key in ['elmk','tmkRemd']:
    for r in status[key]['rows']:
        r['trend']=r['fact']-r['previous'] if r.get('previous') is not None else None

save('mo-data.json',mo);save('operational-mo.json',op);save('mo-details.json',details);save('organization-status.json',status)
print(json.dumps({
 'semd228':{'up':sum((r['trend'] or 0)>0 for r in mo['semd228']['rows']),'down':sum((r['trend'] or 0)<0 for r in mo['semd228']['rows']),'same':sum(r['trend']==0 for r in mo['semd228']['rows'])},
 'smp':{'rows':len(srows),'cards':sum(v['volume'] for v in sdet.values()),'accepted':sum(v['registered'] for v in sdet.values())},
 'tvsp':{k:(sum(r['count'] for r in mo[k]['rows']),sum(v['volume'] for v in details[k].values())) for k in unit_titles}
},ensure_ascii=False,indent=2))
