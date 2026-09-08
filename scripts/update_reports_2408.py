#!/usr/bin/env python3
"""Обновляет три накопительных показателя по срезам 20–21 августа 2026 года."""
import json
import re
from pathlib import Path

from openpyxl import load_workbook

SITE = Path(__file__).resolve().parents[1]
UP = Path('/workspace/scratch/8e713f374d53/upload')


def text(value):
    return '' if value is None else str(value).strip()


def norm(value):
    value = text(value).lower().replace('ё', 'е')
    value = re.sub(r'^(?:филиал\s+)?(?:гауз|гбуз|гбу|фгбу|фгаоу\s*во|ао|ооо|чуз)(?:\s+рт)?\s*', '', value)
    value = re.sub(r'["«»]', '', value)
    value = re.sub(r'\s*\([^)]*\)\s*$', '', value)
    for old, new in {
        'городская клиническая больница': 'гкб',
        'городская детская клиническая больница': 'дгкб',
        'городская детская больница': 'дгб',
        'городская детская поликлиника': 'дгп',
        'детская городская поликлиника': 'дгп',
        'городская больница': 'гб',
        'городская поликлиника': 'гп',
        'центральная районная больница': 'црб',
    }.items():
        value = value.replace(old, new)
    value = re.sub(r'\s*№\s*', '№', value)
    return re.sub(r'[^а-яa-z0-9№]+', ' ', value).strip()


def one(pattern):
    files = sorted(UP.glob(pattern))
    if not files:
        raise FileNotFoundError(pattern)
    return files[-1]


def load(name):
    return json.loads((SITE / 'app' / name).read_text(encoding='utf-8'))


def dump(name, value):
    (SITE / 'app' / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def dataset(old, name, plan, unit, date, period, rows, mode=None, direction=None, note=None):
    if old.get('date') == date:
        previous = {text(row.get('oid')) or norm(row['name']): row.get('previous') for row in old.get('rows', [])}
    else:
        previous = {text(row.get('oid')) or norm(row['name']): row.get('fact') for row in old.get('rows', [])}
    for row in rows:
        prev = previous.get(text(row.get('oid')) or norm(row['name']))
        row['previous'] = prev
        row['trend'] = row['fact'] - prev if prev is not None else None
    result = {'name': name, 'plan': plan, 'unit': unit, 'date': date, 'period': period, 'rows': rows}
    if mode:
        result['mode'] = mode
    if direction:
        result['direction'] = direction
    if note:
        result['note'] = note
    return result


def ratio_rows(path, sheet, name_col, oid_col, denominator_col, numerator_cols):
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    rows, details, oid_details = [], {}, {}
    first = None
    for row_no, row in enumerate(ws.iter_rows(values_only=True), 1):
        name = text(row[name_col]) if len(row) > name_col else ''
        oid = text(row[oid_col]) if len(row) > oid_col else ''
        denominator = row[denominator_col] if len(row) > denominator_col else None
        if not name or not isinstance(denominator, (int, float)):
            continue
        if oid and not re.fullmatch(r'\d+(?:\.\d+)+', oid):
            continue
        if first is None:
            first = row_no
        numerator = sum((row[index] or 0) for index in numerator_cols if len(row) > index and isinstance(row[index], (int, float)))
        fact = numerator / denominator * 100 if denominator else 0
        item = {'name': name, 'oid': oid, 'fact': fact, 'count': int(numerator)}
        rows.append(item)
        detail = {'volume': int(denominator), 'registered': int(numerator)}
        details[norm(name)] = detail
        if oid:
            oid_details[oid] = detail
    if not rows or first is None:
        raise ValueError(f'Не найдены строки данных: {path.name} / {sheet}')
    return rows, details, oid_details, first


mo = load('mo-data.json')
operational = load('operational-mo.json')
details = load('mo-details.json')
detail_oids = load('mo-detail-oids.json')

# 1. Амбулаторный эпикриз: 01.01–21.08.2026.
amb_path = one('Отчет_по_законченному_случаю_амбулаторный*21.08*.xlsx')
amb_rows, amb_details, amb_oid_details, amb_first = ratio_rows(amb_path, '1', 0, 1, 2, [3])
mo['ambulatoryCase'] = dataset(
    mo.get('ambulatoryCase', {}),
    'Доля СЭМД «Эпикриз по законченному случаю амбулаторный» (СЭМД №92, №233) относительно количества случаев',
    95, '%', '21.08.2026', '01.01–21.08.2026', amb_rows,
)
details['ambulatoryCase'], detail_oids['ambulatoryCase'] = amb_details, amb_oid_details

# 2. Госпитализации: лист с единым знаменателем и двумя видами выписных СЭМД.
hosp_path = one('Отчет_по_госпитализациям*21.08*.xlsx')
hosp_rows, hosp_details, hosp_oid_details, hosp_first = ratio_rows(hosp_path, 'Лист3', 1, 2, 3, [4, 5])
mo['hospital'] = dataset(
    mo.get('hospital', {}),
    'Доля СЭМД «Эпикриз в стационаре выписной» и/или «Выписной эпикриз из родильного дома» относительно количества случаев',
    95, '%', '21.08.2026', '01.01–21.08.2026', hosp_rows,
    note='Из знаменателя исключены новорождённые без полиса ОМС. Снижение числа случаев не трактуется как ухудшение МО.',
)
details['hospital'], detail_oids['hospital'] = hosp_details, hosp_oid_details

# Абсолютное количество госпитализаций строится из того же знаменателя.
hosp_count_rows = [
    {'name': row['name'], 'oid': row['oid'], 'fact': hosp_details[norm(row['name'])]['volume'], 'count': hosp_details[norm(row['name'])]['volume']}
    for row in hosp_rows
]
operational['hospitalCount'] = dataset(
    operational.get('hospitalCount', {}), 'Количество госпитализаций', None, '',
    '21.08.2026', '01.01–21.08.2026', hosp_count_rows, mode='count',
    note='Знаменатель показателя выписных эпикризов; новорождённые без полиса ОМС исключены.',
)

# 3. Краткий ввод: амбулаторный и стационарный блоки отдельно, меньше — лучше.
short_path = one('Случаи краткого ввода*20.08*.xlsx')
ws = load_workbook(short_path, read_only=True, data_only=True).active
short_all, short_amb, short_hosp = [], [], []
short_first = None
for row_no, row in enumerate(ws.iter_rows(values_only=True), 1):
    name, oid = text(row[0]), text(row[1])
    if row_no < 7 or not name:
        continue
    if short_first is None:
        short_first = row_no
    ambulatory = int(row[3] or 0) + int(row[6] or 0)
    hospital = int(row[4] or 0) + int(row[5] or 0)
    base = {'name': name, 'oid': oid}
    short_all.append({**base, 'fact': ambulatory + hospital, 'count': ambulatory + hospital})
    short_amb.append({**base, 'fact': ambulatory, 'count': ambulatory})
    short_hosp.append({**base, 'fact': hospital, 'count': hospital})
if not short_all or short_first is None:
    raise ValueError('Не найдены строки краткого ввода')
operational['shortInput'] = dataset(
    operational.get('shortInput', {}), 'Количество случаев краткого ввода — всего', 0, '',
    '20.08.2026', '01.01–20.08.2026', short_all, mode='count', direction='lower',
    note='Накопление является ухудшением. В недельной динамике оценивается прирост новых случаев.',
)
operational['shortInputAmb'] = dataset(
    operational.get('shortInputAmb', {}), 'Краткий ввод — амбулаторный блок и профилактика', 0, '',
    '20.08.2026', '01.01–20.08.2026', short_amb, mode='count', direction='lower',
)
operational['shortInputHosp'] = dataset(
    operational.get('shortInputHosp', {}), 'Краткий ввод — круглосуточный и дневной стационар', 0, '',
    '20.08.2026', '01.01–20.08.2026', short_hosp, mode='count', direction='lower',
)

# Контрольные границы, строки и суммы. При расхождении публикация блокируется.
checks = {
    'ambulatory': {
        'rows': len(amb_rows), 'firstRow': amb_first,
        'numerator': sum(item['registered'] for item in amb_details.values()),
        'denominator': sum(item['volume'] for item in amb_details.values()),
        'first': amb_rows[0]['name'], 'last': amb_rows[-1]['name'],
    },
    'hospital': {
        'rows': len(hosp_rows), 'firstRow': hosp_first,
        'numerator': sum(item['registered'] for item in hosp_details.values()),
        'denominator': sum(item['volume'] for item in hosp_details.values()),
        'first': hosp_rows[0]['name'], 'last': hosp_rows[-1]['name'],
    },
    'shortInput': {
        'rows': len(short_all), 'firstRow': short_first,
        'total': sum(item['fact'] for item in short_all),
        'ambulatory': sum(item['fact'] for item in short_amb),
        'hospital': sum(item['fact'] for item in short_hosp),
        'first': short_all[0]['name'], 'last': short_all[-1]['name'],
    },
}
expected = {
    'ambulatory': (156, 9005261, 10579121),
    'hospital': (107, 426793, 606334),
    'shortInput': (167, 659168, 649689, 9479),
}
actual = {
    'ambulatory': (checks['ambulatory']['rows'], checks['ambulatory']['numerator'], checks['ambulatory']['denominator']),
    'hospital': (checks['hospital']['rows'], checks['hospital']['numerator'], checks['hospital']['denominator']),
    'shortInput': (checks['shortInput']['rows'], checks['shortInput']['total'], checks['shortInput']['ambulatory'], checks['shortInput']['hospital']),
}
if actual != expected:
    raise ValueError(f'Контрольные суммы не сошлись: {actual}')

dump('mo-data.json', mo)
dump('operational-mo.json', operational)
dump('mo-details.json', details)
dump('mo-detail-oids.json', detail_oids)
print(json.dumps(checks, ensure_ascii=False, indent=2))
