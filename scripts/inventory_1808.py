#!/usr/bin/env python3
import json
from pathlib import Path
from openpyxl import load_workbook
import xlrd

ROOT = Path('/workspace/scratch/8e713f374d53/upload')
TOKENS = ('17.08', '18.08')

def val(v):
    if v is None: return None
    s = str(v).strip()
    return s if len(s) <= 160 else s[:157] + '...'

def inspect_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=False)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 14), values_only=True):
            if any(v not in (None, '') for v in row): rows.append([val(v) for v in row[:18]])
        out.append({'sheet': ws.title, 'max_row': ws.max_row, 'max_column': ws.max_column, 'head': rows})
    return out

def inspect_xls(path):
    wb = xlrd.open_workbook(str(path))
    out = []
    for ws in wb.sheets():
        rows = [[val(v) for v in ws.row_values(i)[:18]] for i in range(min(ws.nrows, 14)) if any(v not in ('', None) for v in ws.row_values(i))]
        out.append({'sheet': ws.name, 'max_row': ws.nrows, 'max_column': ws.ncols, 'head': rows})
    return out

def main():
    files = sorted(p for p in ROOT.iterdir() if p.suffix.lower() in ('.xlsx', '.xls') and any(t in p.name for t in TOKENS))
    result = []
    for path in files:
        try:
            sheets = inspect_xlsx(path) if path.suffix.lower() == '.xlsx' else inspect_xls(path)
            result.append({'file': path.name, 'sheets': sheets})
        except Exception as exc:
            result.append({'file': path.name, 'error': repr(exc)})
    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__ == '__main__': main()
