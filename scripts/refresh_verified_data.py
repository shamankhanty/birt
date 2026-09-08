import json
import re
from pathlib import Path

from openpyxl import load_workbook


SITE = Path(__file__).resolve().parents[1]
REVIEW = Path("/workspace/scratch/8e713f374d53/review")


def dump(name, value):
    (SITE / "app" / name).write_text(
        json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def find_calc(fragment):
    return next(
        p
        for p in REVIEW.rglob("*.xlsx")
        if "Расчеты" in str(p) and fragment in p.name
    )


def key(name):
    value = re.sub(r'^(?:Филиал\s+)?(?:ГАУЗ|ГБУЗ|ГБУ|ФГБУ|ФГАОУВО|ФГАОУ ВО|АО|ООО)(?:\s+РТ)?\s*', '', name, flags=re.I)
    value = re.sub(r'["«»]', '', value)
    value = re.sub(r'\s*\([^)]*\)\s*$', '', value)
    return re.sub(r'\s+', ' ', value).strip().lower()


def tvsp_rows(path, previous_path, mode="lists"):
    def read_one(book_path):
        wb = load_workbook(book_path, read_only=True, data_only=True)
        if mode == "lists":
            ws = next(w for w in wb.worksheets if "Детализа" in w.title)
            rows = [r for r in ws.iter_rows(min_row=7, values_only=True) if r[0] == "Республика Татарстан" and r[1]]
            name_i, fact_i = 1, 5
            is_fact = lambda r: str(r[fact_i]).strip().lower() != "нет"
        else:
            ws = next(w for w in wb.worksheets if "Факт передачи" in w.title)
            rows = [r for r in ws.iter_rows(min_row=7 if mode == "diagnostic" else 6, values_only=True) if len(r) > 6 and r[1] == "Республика Татарстан" and r[2]]
            name_i = 2
            is_fact = lambda r: str(r[6]).strip().lower() == "да"
        result = {}
        for row in rows:
            item = result.setdefault(key(str(row[name_i])), {"name": str(row[name_i]), "volume": 0, "registered": 0})
            item["volume"] += 1
            item["registered"] += int(is_fact(row))
        return result

    current, previous = read_one(path), read_one(previous_path)
    rows, details = [], {}
    for k, item in current.items():
        fact = item["registered"] / item["volume"] * 100 if item["volume"] else 0
        old = previous.get(k)
        previous_fact = old["registered"] / old["volume"] * 100 if old and old["volume"] else None
        rows.append({
            "name": item["name"],
            "fact": fact,
            "previous": previous_fact,
            "trend": fact - previous_fact if previous_fact is not None else None,
        })
        details[k] = {"volume": item["volume"], "registered": item["registered"]}
    return rows, details


def ambulatory_case():
    path = find_calc("10, 11 сдайды Отчет_по_законченному")
    wb = load_workbook(path, read_only=True, data_only=True)
    current, previous = wb["01.01.-29.07."], wb["01.01.-21.07."]
    old = {}
    for row in previous.iter_rows(min_row=5, values_only=True):
        if row[0] and isinstance(row[2], (int, float)) and isinstance(row[3], (int, float)):
            old[key(str(row[0]))] = row[3] / row[2] * 100
    rows, details = [], {}
    for row in current.iter_rows(min_row=5, values_only=True):
        if not row[0] or str(row[0]).lower().startswith("итого") or not isinstance(row[2], (int, float)) or not isinstance(row[3], (int, float)):
            continue
        fact = row[3] / row[2] * 100
        prev = old.get(key(str(row[0])))
        rows.append({"name": str(row[0]), "fact": fact, "previous": prev, "trend": fact - prev if prev is not None else None})
        details[key(str(row[0]))] = {"volume": row[2], "registered": row[3]}
    return rows, details


mo_path = SITE / "app" / "mo-data.json"
details_path = SITE / "app" / "mo-details.json"
mo = json.loads(mo_path.read_text(encoding="utf-8"))
details = json.loads(details_path.read_text(encoding="utf-8"))

tvsp_specs = {
    "tvspStationary": ("9 слайд Отчет_Доля_ТВСП", "15 слайд Отчет_Доля_ТВСП", "lists", "ТВСП: выписной эпикриз стационара и родильного дома"),
    "tvspAmbulatory": ("11 слайд Отчет_Доля_ТВСП", "14 слайд Отчет_Доля_ТВСП", "lists", "ТВСП: амбулаторный эпикриз, талон пациента и протокол консультации"),
    "tvspLaboratory": ("10 слайд Отчет_Доля_ТВСП", "14 сдайд Отчет_Доля_ТВСП", "lists", "ТВСП: протокол лабораторного исследования"),
    "tvspDiagnostic": ("10 слайд Доля_ТВСП", "14 слайд Доля_ТВСП", "diagnostic", "ТВСП: протокол диагностических исследований"),
    "smpFederal": ("11 слайд Отчет_СМП_ТВСП", "14 слайд Отчет_СМП_ТВСП", "smp", "Станции СМП, передающие карты вызова в РЭМД"),
}

for metric, (cur, prev, mode, name) in tvsp_specs.items():
    rows, metric_details = tvsp_rows(find_calc(cur), find_calc(prev), mode)
    mo[metric] = {"name": name, "plan": 100, "unit": "%", "date": "06.08.2026", "rows": rows}
    details[metric] = metric_details

mo["tvspLaboratory"]["note"] = "Итог по РТ — 97,50%; детализация по МО показана по федеральной выгрузке"

rows, metric_details = ambulatory_case()
mo["ambulatoryCase"] = {
    "name": "Эпикриз по законченному амбулаторному случаю",
    "plan": 95,
    "unit": "%",
    "date": "29.07.2026",
    "rows": rows,
    "note": "Последний достоверный накопительный срез: 01.01–29.07.2026",
}
details["ambulatoryCase"] = metric_details

dump("mo-data.json", mo)
dump("mo-details.json", details)
