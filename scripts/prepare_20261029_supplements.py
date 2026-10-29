"""Prepare independently sourced views in an unpublished staging app."""
import argparse, hashlib, json, re, sys
from collections import defaultdict
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts/pipeline'))
from adapters import parse_max,source_rows
from source_catalog import scan
from transactional_stage import validate

def read(p):return json.loads(p.read_text(encoding='utf-8'))
def write(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')
def prepare(source_folder,stage_folder):
 app=stage_folder/'app'; manifest=read(stage_folder/'staging-manifest.json'); before={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in app.glob('*.json')}
 sources=list(source_folder.rglob('*.xlsx'))
 remd=next(p for p in sources if p.name.startswith('1.Отчет СЭМД_РЭМД'))
 wb=openpyxl.load_workbook(remd,read_only=True,data_only=True);rows=list(wb['Отчет РЭМД'].values);wb.close()
 dates=re.findall(r'\d{2}\.\d{2}\.\d{4}',str(rows[0][0]));assert len(dates)==2
 names,formats,values=rows[4],rows[5],rows[7];grouped=defaultdict(lambda:{'count':0,'formats':[]});last=''
 for i in range(3,len(values)):
  if names[i]:last=str(names[i]).strip()
  if not last or (values[i] is None and names[i] is None and formats[i] is None):continue
  value=values[i];assert isinstance(value,(int,float)) and value>=0,(i,last,value)
  grouped[last]['count']+=int(value)
  if formats[i] and str(formats[i]) not in grouped[last]['formats']:grouped[last]['formats'].append(str(formats[i]))
 summary={'period':'–'.join(dates),'formed':str(rows[1][0]).split(': ',1)[-1],'source':remd.name,'total':int(values[2]),'registeredTypes':int(values[1]),'items':[{'name':k,'count':v['count'],'format':' / '.join(v['formats']) or '—'} for k,v in grouped.items()]}
 assert sum(x['count'] for x in summary['items'])==summary['total'];assert sum(x['count']>0 for x in summary['items'])==summary['registeredTypes'];write(app/'semd-summary.json',summary)
 control=read(app/'federal-control.json')
 for point,needle in [('37','Экстренное извещение об инфекционном заболевании'),('38','Выписка из протокола решения врачебной комиссии'),('39','Льготный рецепт')]:
  found=[x for x in summary['items'] if (x['name'].casefold()==needle.casefold() if point=='38' else x['name'].casefold().startswith(needle.casefold()))];assert len(found)==1,(point,found)
  row=next(x for x in control['agreement'] if x['id']==point)
  row.update(regionalFact=f"{found[0]['count']:,} СЭМД".replace(',','\u00a0'),regionalPeriod=f"01.01–{dates[1]}",sourceNote=f"Источник: {remd.name}. Количество зарегистрированных СЭМД. Автоматизированный способ формирования отдельно не верифицировался.")
 write(app/'federal-control.json',control)
 # Preserve the preceding accepted monthly and cumulative slices independently.
 views=read(app/'max-period-views.json')
 for kind,start in [('cumulative','01.01.'),('operational','01.10.')]:
  source=next(p for p in sources if p.name.startswith('ТМК_МАХ '+start))
  book=openpyxl.load_workbook(source,read_only=True,data_only=True);heading=str(book.active.cell(1,1).value);book.close()
  sd,ed=re.findall(r'\d{2}\.\d{2}\.\d{4}',heading);old=views[kind]
  entry={**old,'date':ed,'period':f'{sd[:5]}–{ed}','source':source.name,'previousDate':old['date'],'previousPeriod':old['period'],'previousSource':old['source'],'metrics':{}}
  for service,metric,col in [('tmk','tmkMaxCount',3),('eln','elnMaxCount',4)]:
   raw=parse_max(source,col);indexed={r.get('oid') or r['name']:r for r in source_rows(raw,metric)}
   oldmetric=old['metrics'][service];current=[]
   for row in oldmetric['rows']:
    found=indexed.get(row['oid']);current.append({**row,'fact':found['fact'] if found else None,'previous':row['fact']})
   entry['metrics'][service]={**oldmetric,'sourceTotal':sum(r['count'] for r in raw),'previousSourceTotal':oldmetric['sourceTotal'],'plannedMoFact':sum(r['fact'] for r in current if r['fact'] is not None),'previousPlannedMoFact':oldmetric['plannedMoFact'],'rows':current}
   assert len(current)==85
  views[kind]=entry
 write(app/'max-period-views.json',views)
 validation=validate(app,stage_folder/'supplement-validation');assert validation['overallStatus']!='FAIL',validation['summary']
 changed=[p.name for p in app.glob('*.json') if before.get(p.name)!=hashlib.sha256(p.read_bytes()).hexdigest()]
 manifest['changedFiles']=sorted(set(manifest['changedFiles']+changed));manifest['formalValidation']={'status':validation['overallStatus'],'summary':validation['summary']};write(stage_folder/'staging-manifest.json',manifest)
 evidence={'sourceFiles':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},'statistics':manifest['statistics'],'retained':[{'metric':x['metric'],'errors':x['errors']}for x in manifest['indicators']if x['state']=='RETAINED'],'supplementalFiles':changed,'semdSummary':{k:summary[k] for k in ['period','total','registeredTypes']},'maxViews':{kind:{'date':views[kind]['date'],'previousDate':views[kind]['previousDate'],'metrics':{k:{'sourceTotal':v['sourceTotal'],'plannedMoFact':v['plannedMoFact']} for k,v in views[kind]['metrics'].items()}}for kind in ['cumulative','operational']},'monthlyRatingChanged':False,'published':False}
 write(stage_folder/'update-evidence.json',evidence);print(json.dumps(evidence,ensure_ascii=False,indent=2))
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--input',type=Path,required=True);p.add_argument('--stage',type=Path,required=True);a=p.parse_args();prepare(a.input,a.stage)
