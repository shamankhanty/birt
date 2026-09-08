#!/usr/bin/env python3
"""Load the combined DOGVN/PMO + SEMD 122/228 report dated 01.09.2026."""
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/workspace/scratch/1f5740e14d75/upload/Закрытые случаи ДОГВН ПМО и зарегистрированные СЭМД 122 и 228.xlsx")


def clean(value):
    return re.sub(r"\s+", " ", str(value or "").strip())


def short_name(value):
    value = clean(value)
    value = re.sub(r"^(?:ГАУЗ|ГБУЗ|ФГАОУВО|ЧУЗ)\s*", "", value, flags=re.I)
    return value.strip(' «»"')


def main():
    ws = load_workbook(SOURCE, read_only=True, data_only=True).active
    rows = []
    for values in ws.iter_rows(min_row=5, values_only=True):
        name = clean(values[1])
        if not name:
            continue
        oid = clean(values[2])
        dogvn = int(values[3] or 0)
        pmo = int(values[4] or 0)
        semd122_raw = values[9]
        semd228_raw = values[13]
        # In this combined report the source columns are present for every row;
        # an empty status cell is the report's representation of no registrations.
        semd122 = int(semd122_raw or 0)
        semd228 = int(semd228_raw or 0)
        selected = max(semd122, semd228)
        selected_type = "122" if semd122 >= semd228 else "228"
        foms = dogvn + pmo
        share = selected / foms * 100 if foms else None
        issues = []
        if semd122_raw is None:
            issues.append("Пустое поле СЭМД 122 в источнике принято как 0")
        if semd228_raw is None:
            issues.append("Пустое поле СЭМД 228 в источнике принято как 0")
        if foms == 0 and selected:
            issues.append("СЭМД есть, знаменатель равен 0")
        if share is not None and share > 100:
            issues.append("Значение выше 100%")
        rows.append({
            "name": name, "oid": oid, "child": False,
            "semd122": semd122, "semd228": semd228,
            "selected": selected, "selectedType": selected_type,
            "foms": foms, "share": share, "oldShare": None,
            "change": None, "issues": issues,
        })

    numerator = sum(row["selected"] for row in rows)
    denominator = sum(row["foms"] for row in rows)
    summary = {
        "status": "ready", "year": 2026,
        "formula": "MAX(СЭМД 122; СЭМД 228)",
        "period122": "01.01.2026–28.08.2026",
        "period228": "01.01.2026–28.08.2026",
        "source122": SOURCE.name, "source228": SOURCE.name,
        "organizations": len(rows), "changed": 0, "changedChildren": 0,
        "over100": sum(bool(row["share"] and row["share"] > 100) for row in rows),
        "zeroDenominatorWithSemd": sum(row["foms"] == 0 and row["selected"] > 0 for row in rows),
        "missing": sum(bool(row["issues"]) for row in rows),
        "childrenMissing": True, "childOrganizations": 0,
        "numerator": numerator, "denominator": denominator,
        "share": numerator / denominator * 100,
        "selected122": sum(row["semd122"] > row["semd228"] for row in rows),
        "selected228": sum(row["semd228"] > row["semd122"] for row in rows),
        "selectedEqual": sum(row["semd122"] == row["semd228"] for row in rows),
    }
    (ROOT / "app/preventive-semd-audit.json").write_text(
        json.dumps({"summary": summary, "rows": rows}, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    mo_data_path = ROOT / "app/mo-data.json"
    mo_data = json.loads(mo_data_path.read_text(encoding="utf-8"))
    mo_data["semd228"] = {
        "name": "Доля СЭМД (122/228) профилактического осмотра (диспансеризации)",
        "plan": 95, "unit": "%", "date": "28.08.2026", "period": "01.01–28.08.2026",
        "note": "Расчёт 2026 года: MAX(СЭМД 122; СЭМД 228) отдельно по каждой МО. В источнике 78 взрослых МО; детские МО отсутствуют и не оцениваются до перевыгрузки. Пустые ячейки регистрации в представленном комбинированном отчёте трактуются как отсутствие зарегистрированных документов.",
        "rows": [{
            "name": row["name"], "oid": row["oid"], "fact": row["share"] or 0,
            "count": row["selected"], "previous": None, "trend": None,
            "sourceWarning": "Динамика не рассчитывается: состав МО и знаменатель отличаются от предыдущей выгрузки."
        } for row in rows],
    }
    mo_data_path.write_text(json.dumps(mo_data, ensure_ascii=False, indent=2), encoding="utf-8")

    details_path = ROOT / "app/mo-details.json"
    details = json.loads(details_path.read_text(encoding="utf-8"))
    details["semd228"] = {short_name(row["name"]).casefold(): {"volume": row["foms"], "registered": row["selected"]} for row in rows}
    details_path.write_text(json.dumps(details, ensure_ascii=False, indent=2), encoding="utf-8")

    oid_path = ROOT / "app/mo-detail-oids.json"
    oid_details = json.loads(oid_path.read_text(encoding="utf-8"))
    oid_details["semd228"] = {row["oid"]: {"volume": row["foms"], "registered": row["selected"]} for row in rows}
    oid_path.write_text(json.dumps(oid_details, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
