"""Read immutable original September exports; never mutate operational datasets."""
import hashlib,json,sys
from pathlib import Path
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'pipeline'))
from adapters import parse_remd,parse_foms
root=Path(sys.argv[1]); dest=Path(sys.argv[2])
def source(pattern):
 p=next(root.glob(pattern));return p,{'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
r,rs=source('*СЭМД_РЭМД*.xlsx');f,fs=source('Количество_СЭМД*.xlsx');h,hs=source('Отчет по госпитализациям*.xlsx')
e=next(root.rglob('*ЕГПУ*01.09.26-30.09.26*.xlsx'));es={'file':e.name,'sha256':hashlib.sha256(e.read_bytes()).hexdigest()}
rr,ff=parse_remd(r),parse_foms(f); metrics={'semd228':[],'hospital':[],'egpu':[],'egpu2days':[]}
for oid,v in ff.items():
 a=rr.get(oid);metrics['semd228'].append({'oid':oid,'name':v['name'],'numerator':max(a['s122'],a['s228']) if a else None,'denominator':v['den'],'s122':a['s122'] if a else None,'s228':a['s228'] if a else None})
for v in openpyxl.load_workbook(h,read_only=True,data_only=True).active.iter_rows(min_row=6,values_only=True):
 if v[2] and str(v[2]).startswith('1.2.'):
  metrics['hospital'].append({'oid':str(v[2]).strip(),'name':v[1],'numerator':(v[4] or 0)+(v[5] or 0),'denominator':v[3]})
for v in openpyxl.load_workbook(e,read_only=True,data_only=True).active.iter_rows(min_row=4,values_only=True):
 if v[3] and str(v[3]).startswith('1.2.'):
  for id,col in [('egpu',5),('egpu2days',6)]:metrics[id].append({'oid':str(v[3]).strip(),'name':v[2],'numerator':v[col],'denominator':v[4]})
for id,rows in metrics.items():
 assert len({x['oid'] for x in rows})==len(rows),id
 for x in rows:
  for k in ['numerator','denominator']:assert x[k] is None or isinstance(x[k],(int,float)),(id,x)
assert sum(x['numerator'] for x in metrics['egpu'])==6684
assert sum(x['denominator'] for x in metrics['egpu'])==6690
assert sum(x['numerator'] for x in metrics['egpu2days'])==6077
out={'sources':{'remd':rs,'foms':fs,'hospital':hs,'egpu':es},'metrics':metrics}
dest.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
print({k:len(v) for k,v in metrics.items()})
