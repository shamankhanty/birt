#!/usr/bin/env python3
"""Пересчёт свидетельств: числитель — только «Зарегистрирован»."""
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

SITE = Path(__file__).resolve().parents[1]
UPLOAD = Path('/workspace/scratch/8e713f374d53/upload')

CURRENT = {
    'birth': next(UPLOAD.glob('Свидетельства_о_рождении*01.01.-17.08*.xlsx')),
    'death': next(UPLOAD.glob('I_Свид-ва о смерти*01.01.-17.08*.xlsx')),
}
PREVIOUS = {
    'birth': next(UPLOAD.glob('Свидетельства_о_рождении*01.01.-10.08*.xlsx')),
    'death': next(UPLOAD.glob('I_Свид-ва о смерти*01.01.-10.08*.xlsx')),
}

def text(value): return '' if value is None else str(value).strip()
def norm(value):
    value=text(value).lower().replace('ё','е')
    value=re.sub(r'^(?:филиал\s+)?(?:гауз|гбуз|гбу|фгбу|фгаоу\s*во|ао|ооо|чуз)(?:\s+рт)?\s*','',value)
    value=re.sub(r'["«»]','',value)
    value=re.sub(r'\s*\([^)]*\)\s*$','',value)
    return re.sub(r'[^а-яa-z0-9№]+',' ',value).strip()

def read(path):
    ws=load_workbook(path,read_only=True,data_only=True).active
    by_org=defaultdict(lambda:{'total':0,'registered':0,'statuses':Counter()})
    seen=set()
    first=[]; last=[]
    for row in ws.iter_rows(values_only=True):
        name=text(row[0]) if row else ''
        number=text(row[1]) if len(row)>1 else ''
        state=text(row[3]) if len(row)>3 else ''
        if not name or not number or name.startswith(('Учреждение','Где в столбце')): continue
        if number in seen: raise ValueError(f'Дубль свидетельства {number}: {path.name}')
        seen.add(number)
        item=by_org[name]; item['total']+=1; item['registered']+=int(state.lower()=='зарегистрирован'); item['statuses'][state or 'Пусто']+=1
        if len(first)<3: first.append((name,number,state))
        last=(last+[(name,number,state)])[-3:]
    return by_org,seen,first,last

mo=json.loads((SITE/'app/mo-data.json').read_text())
details=json.loads((SITE/'app/mo-details.json').read_text())
summary={}
for metric in ('birth','death'):
    current,seen,first,last=read(CURRENT[metric]); previous,_,_,_=read(PREVIOUS[metric])
    rows=[]; metric_details={}; statuses=Counter()
    previous_by_key={norm(name):value for name,value in previous.items()}
    for name,value in current.items():
        total=value['total']; registered=value['registered']; fact=registered/total*100 if total else 0
        old=previous_by_key.get(norm(name)); old_fact=old['registered']/old['total']*100 if old and old['total'] else None
        rows.append({'name':name,'fact':fact,'count':registered,'previous':old_fact,'trend':fact-old_fact if old_fact is not None else None})
        metric_details[norm(name)]={'volume':total,'registered':registered}
        statuses.update(value['statuses'])
    rows.sort(key=lambda row:norm(row['name']))
    mo[metric].update({'date':'17.08.2026','period':'01.01–17.08.2026','rows':rows})
    mo[metric]['note']='В знаменатель включены все выданные свидетельства; в числитель — только статус «Зарегистрирован». Статусы «Создан», «Отправлен», «Ошибка» и пустой статус требуют контроля.'
    details[metric]=metric_details
    total=sum(v['total'] for v in current.values()); registered=sum(v['registered'] for v in current.values())
    summary[metric]={'total':total,'registered':registered,'fact':registered/total*100,'organizations':len(current),'attention':sum(v['registered']/v['total']*100<99 for v in current.values() if v['total']),'statuses':dict(statuses),'first':first,'last':last,'unique':len(seen)}

(SITE/'app/mo-data.json').write_text(json.dumps(mo,ensure_ascii=False,indent=2))
(SITE/'app/mo-details.json').write_text(json.dumps(details,ensure_ascii=False,indent=2))
print(json.dumps(summary,ensure_ascii=False,indent=2))
