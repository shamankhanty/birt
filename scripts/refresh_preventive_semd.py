#!/usr/bin/env python3
"""Recalculate the preventive-examination indicator from SЭМД 122/228.

For 2026, the numerator is MAX(122, 228) for each organization and period.
For 2027+, only SЭМД 228 is used. Missing source files or organization rows are
kept as missing and never converted to zero.
"""
import argparse
import csv
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def text(value):
    return re.sub(r"\s+", " ", str(value or "").strip())


def find_header(ws, required):
    for number, row in enumerate(ws.iter_rows(values_only=True), 1):
        cells = [text(value).casefold() for value in row]
        if all(any(token in cell for cell in cells) for token in required):
            return number, cells
    raise ValueError(f"Не найдена строка заголовков: {required}")


def column(cells, variants):
    for index, cell in enumerate(cells):
        if all(part in cell for part in variants):
            return index
    raise ValueError(f"Не найден столбец: {variants}")


def period_from_book(ws):
    heading = " ".join(text(ws.cell(row=i, column=j).value) for i in range(1, min(6, ws.max_row) + 1) for j in range(1, min(8, ws.max_column) + 1))
    match = re.search(r"(?:с\s*)?(\d{2}\.\d{2}\.\d{4})\s*(?:по|[-–])\s*(\d{2}\.\d{2}\.\d{4})", heading)
    return f"{match.group(1)}–{match.group(2)}" if match else None


def read_228(path):
    ws = load_workbook(path, read_only=True, data_only=True).active
    header, cells = find_header(ws, ["наименование", "oid", "обращен", "228"])
    name_col = column(cells, ["наименование"]); oid_col = column(cells, ["oid"])
    denominator_col = column(cells, ["обращен"]); count_col = column(cells, ["228"])
    rows = {}
    for values in ws.iter_rows(min_row=header + 1, values_only=True):
        name = text(values[name_col]); oid = text(values[oid_col])
        if not name or name.casefold().startswith("итого"): continue
        rows[oid or f"name:{name.casefold()}"] = {"name": name, "oid": oid, "semd228": int(values[count_col]) if isinstance(values[count_col], (int, float)) else None, "foms": int(values[denominator_col]) if isinstance(values[denominator_col], (int, float)) else None}
    return rows, period_from_book(ws)


def read_122(path):
    ws = load_workbook(path, read_only=True, data_only=True).active
    header, cells = find_header(ws, ["наименование", "oid", "122"])
    name_col = column(cells, ["наименование"]); oid_col = column(cells, ["oid"]); count_col = column(cells, ["122"])
    rows = {}
    for values in ws.iter_rows(min_row=header + 1, values_only=True):
        name = text(values[name_col]); oid = text(values[oid_col])
        if not name or name.casefold().startswith("итого"): continue
        rows[oid or f"name:{name.casefold()}"] = {"name": name, "oid": oid, "semd122": int(values[count_col]) if isinstance(values[count_col], (int, float)) else None}
    return rows, period_from_book(ws)


def child_oids():
    registry = json.loads((ROOT / "app/mo-registry.json").read_text(encoding="utf-8"))
    return {row["oid"] for row in registry["organizations"] if re.search(r"детск|\b(?:адрб|дркб|ндрб|кдмц)\b", (row.get("type", "") + " " + row.get("name", "")).casefold())}


def build_rows(source_228, source_122, year, old_details):
    children = child_oids(); keys = sorted(set(source_228) | set(source_122 or {}))
    result = []
    for key in keys:
        a = (source_122 or {}).get(key); b = source_228.get(key)
        semd122 = a.get("semd122") if a else None; semd228 = b.get("semd228") if b else None; foms = b.get("foms") if b else None
        if year <= 2026:
            selected = max(semd122, semd228) if semd122 is not None and semd228 is not None else None
            selected_type = "122" if selected is not None and semd122 >= semd228 else "228" if selected is not None else None
        else:
            selected, selected_type = semd228, "228" if semd228 is not None else None
        share = selected / foms * 100 if selected is not None and foms not in (None, 0) else None
        name = (b or a or {}).get("name", key); oid = (b or a or {}).get("oid", "")
        old = old_details.get(oid) or old_details.get(name.casefold()) or {}
        old_share = old.get("registered", 0) / old["volume"] * 100 if old.get("volume") else None
        issues=[]
        if semd122 is None and year <= 2026: issues.append("Нет данных СЭМД 122")
        if semd228 is None: issues.append("Нет данных СЭМД 228")
        if foms is None: issues.append("Нет данных ФОМС")
        if foms == 0 and (semd122 or semd228): issues.append("СЭМД есть, знаменатель равен 0")
        if share is not None and share > 100: issues.append("Значение выше 100%")
        result.append({"name":name,"oid":oid,"child":oid in children,"semd122":semd122,"semd228":semd228,"selected":selected,"selectedType":selected_type,"foms":foms,"share":share,"oldShare":old_share,"change":share-old_share if share is not None and old_share is not None else None,"issues":issues})
    return result


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--semd228", type=Path, required=True)
    parser.add_argument("--semd122", type=Path)
    parser.add_argument("--year", type=int, default=2026)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--csv", type=Path, default=ROOT / "deliverables/Проверка_СЭМД_122_228_2026.csv")
    args=parser.parse_args()
    source_228, period_228 = read_228(args.semd228)
    source_122, period_122 = read_122(args.semd122) if args.semd122 else (None, None)
    details=json.loads((ROOT/"app/mo-detail-oids.json").read_text(encoding="utf-8")).get("semd228",{})
    by_name=json.loads((ROOT/"app/mo-details.json").read_text(encoding="utf-8")).get("semd228",{})
    old={**by_name,**details}
    rows=build_rows(source_228,source_122,args.year,old)
    comparable = bool(source_122) if args.year <= 2026 else True
    same_period = period_122 == period_228 if args.year <= 2026 and period_122 else args.year > 2026
    status="ready" if comparable and same_period else "blocked"
    summary={"status":status,"year":args.year,"formula":"MAX(СЭМД 122; СЭМД 228)" if args.year<=2026 else "СЭМД 228","period122":period_122,"period228":period_228,"source122":args.semd122.name if args.semd122 else None,"source228":args.semd228.name,"organizations":len(rows),"changed":sum(r["change"] not in (None,0) for r in rows),"changedChildren":sum(r["child"] and r["change"] not in (None,0) for r in rows),"over100":sum(r["share"] is not None and r["share"]>100 for r in rows),"zeroDenominatorWithSemd":sum(r["foms"]==0 and bool((r["semd122"] or 0)+(r["semd228"] or 0)) for r in rows),"missing":sum(bool(r["issues"]) for r in rows)}
    payload={"summary":summary,"rows":rows}
    (ROOT/"app/preventive-semd-audit.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    args.csv.parent.mkdir(parents=True,exist_ok=True)
    with args.csv.open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f,delimiter=";"); writer.writerow(["МО","Детская МО","СЭМД 122","СЭМД 228","Выбрано в числитель","Вид","Случаи ФОМС","Доля после, %","Доля до, %","Изменение, п.п.","Ошибки"])
        for r in rows: writer.writerow([r["name"],"Да" if r["child"] else "Нет",r["semd122"],r["semd228"],r["selected"],r["selectedType"],r["foms"],r["share"],r["oldShare"],r["change"],"; ".join(r["issues"])])
    if args.apply:
        if status != "ready": raise SystemExit("Расчёт заблокирован: отсутствует СЭМД 122 или периоды источников не совпадают")
        raise SystemExit("Применение к рабочим JSON будет разрешено после контрольной сверки итогов; используйте сформированную таблицу")
    print(json.dumps(summary,ensure_ascii=False,indent=2))


if __name__ == "__main__": main()
