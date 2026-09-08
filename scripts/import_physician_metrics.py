#!/usr/bin/env python3
"""Build verified monthly physician metrics from the federal BI workbook."""

from __future__ import annotations

import json
import math
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl


SPECIALTIES = {
    "Акушер-гинеколог": ("doctor500_obgyn", "Доля врачей акушеров-гинекологов ПМСП, подписавших свыше 500 СЭМД"),
    "Врач общей практики": ("doctor500_gp", "Доля врачей общей практики, подписавших свыше 500 СЭМД"),
    "Кардиолог": ("doctor500_cardiologist", "Доля врачей-кардиологов ПМСП, подписавших свыше 500 СЭМД"),
    "Онколог": ("doctor500_oncologist", "Доля врачей-онкологов ПМСП, подписавших свыше 500 СЭМД"),
    "Офтальмолог": ("doctor500_ophthalmologist", "Доля врачей-офтальмологов ПМСП, подписавших свыше 500 СЭМД"),
    "Педиатр": ("doctor500_pediatrician", "Доля врачей-педиатров, подписавших свыше 500 СЭМД"),
    "Стоматолог": ("doctor500_dentist", "Доля врачей-стоматологов ПМСП, подписавших свыше 500 СЭМД"),
    "Терапевт": ("doctor500_therapist", "Доля врачей-терапевтов, подписавших свыше 500 СЭМД"),
    "Хирург": ("doctor500_surgeon", "Доля врачей-хирургов ПМСП, подписавших свыше 500 СЭМД"),
}


def share(num: int, den: int) -> float:
    return num / den * 100 if den else 0.0


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("usage: import_physician_metrics.py INPUT_XLSX REGISTRY_JSON OUTPUT_JSON")
    source, registry_path, output = map(Path, sys.argv[1:])
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    registry = json.loads(registry_path.read_text(encoding="utf-8"))["organizations"]
    registry_oids = {row["oid"] for row in registry}

    all_ws = wb["Все врачи_Детализация по МО"]
    all_rows = [r for r in all_ws.iter_rows(min_row=9, values_only=True) if r[0] == "Республика Татарстан" and r[1]]
    specialty_ws = wb["Врачи по спец-тям_По МО"]
    specialty_rows = [r for r in specialty_ws.iter_rows(min_row=9, values_only=True) if r[0] == "Республика Татарстан" and r[1] and r[4] in SPECIALTIES]

    all_oids = {str(r[1]) for r in all_rows}
    specialty_oids = {str(r[1]) for r in specialty_rows}
    if all_oids - registry_oids or specialty_oids - registry_oids:
        raise ValueError(f"Unmatched OIDs: all={sorted(all_oids-registry_oids)}, specialty={sorted(specialty_oids-registry_oids)}")
    if len(all_rows) != 125 or len(all_oids) != 125:
        raise ValueError(f"Expected 125 unique all-doctor MO rows, got {len(all_rows)} rows/{len(all_oids)} OIDs")
    if len(specialty_rows) != 726 or len(specialty_oids) != 116:
        raise ValueError(f"Expected 726 specialty rows and 116 OIDs, got {len(specialty_rows)}/{len(specialty_oids)}")

    for r in specialty_rows:
        denominator, signed, under100, between, over500 = map(int, r[5:10])
        if signed != under100 + between + over500 or signed > denominator:
            raise ValueError(f"Category reconciliation failed: {r[1]} {r[4]}")
        if not math.isclose(float(r[13]), share(over500, denominator), abs_tol=0.011):
            raise ValueError(f"Share mismatch: {r[1]} {r[4]}")

    datasets: dict[str, dict] = {}

    def make_rows(rows, numerator_index, denominator_index, small_rule=False):
        prepared = []
        for r in rows:
            denominator = int(r[denominator_index])
            numerator = int(r[numerator_index])
            prepared.append({
                "name": str(r[2]), "oid": str(r[1]), "fact": share(numerator, denominator),
                "count": numerator, "volume": denominator, "previous": None, "trend": None,
                "sourceWarning": "Справочно: в МО только 1–2 врача этой специальности; в заслушивании не оценивается." if small_rule and denominator < 3 else None,
            })
        return prepared

    all_num = sum(int(r[6]) for r in all_rows)
    all_den = sum(int(r[4]) for r in all_rows)
    datasets["doctorsAll"] = {
        "name": "Доля врачей, подписавших хотя бы один СЭМД", "plan": 80, "unit": "%",
        "date": "31.08.2026", "period": "август 2026", "periodType": "month",
        "note": "Месячный результат. Числитель — врачи, подписавшие хотя бы один СЭМД; знаменатель — врачи по ФРМР.",
        "summary": {"numerator": all_num, "denominator": all_den, "fact": share(all_num, all_den)},
        "rows": make_rows(all_rows, 6, 4),
    }
    level3 = [r for r in all_rows if r[3] == "III уровень"]
    level3_num = sum(int(r[6]) for r in level3)
    level3_den = sum(int(r[4]) for r in level3)
    datasets["doctorsLevel3"] = {
        "name": "Доля врачей МО III уровня, подписавших хотя бы один СЭМД", "plan": 90, "unit": "%",
        "date": "31.08.2026", "period": "август 2026", "periodType": "month",
        "note": "Месячный результат по МО III уровня.",
        "summary": {"numerator": level3_num, "denominator": level3_den, "fact": share(level3_num, level3_den)},
        "rows": make_rows(level3, 6, 4),
    }

    by_specialty = defaultdict(list)
    for row in specialty_rows:
        by_specialty[row[4]].append(row)
    for specialty, (metric_id, title) in SPECIALTIES.items():
        rows = by_specialty[specialty]
        numerator = sum(int(r[9]) for r in rows)
        denominator = sum(int(r[5]) for r in rows)
        datasets[metric_id] = {
            "name": title, "plan": 50, "unit": "%", "date": "31.08.2026",
            "period": "август 2026", "periodType": "month", "specialty": specialty,
            "note": "Месячный результат. При 1–2 врачах показатель показывается справочно и не влияет на приоритет заслушивания.",
            "summary": {"numerator": numerator, "denominator": denominator, "fact": share(numerator, denominator)},
            "rows": make_rows(rows, 9, 5, small_rule=True),
        }

    subject = wb["Врачи по спец-тям_По субъекту"]
    subject_values = {str(r[1]): (int(r[2]), int(r[6])) for r in subject.iter_rows(min_row=10, values_only=True) if r[0] == "Республика Татарстан"}
    for specialty, (metric_id, _) in SPECIALTIES.items():
        expected = subject_values[specialty]
        actual = (datasets[metric_id]["summary"]["denominator"], datasets[metric_id]["summary"]["numerator"])
        if expected != actual:
            raise ValueError(f"Regional reconciliation failed for {specialty}: {expected} != {actual}")

    payload = {
        "source": source.name, "formed": "31.08.2026", "period": "август 2026",
        "quality": {"allMoRows": 125, "specialtyRows": 726, "allMoOids": 125, "specialtyMoOids": 116, "unmatchedOids": 0, "categoryErrors": 0},
        "datasets": datasets,
    }
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"datasets": len(datasets), **payload["quality"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
