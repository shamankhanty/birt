#!/usr/bin/env python3
import json,re
from pathlib import Path
from openpyxl import load_workbook

SITE=Path(__file__).resolve().parents[1];UP=Path('/workspace/scratch/8e713f374d53/upload')
def text(v):return '' if v is None else str(v).strip()
def one(pattern):
    fs=list(UP.glob(pattern))
    if not fs:raise FileNotFoundError(pattern)
    return fs[-1]
def oids(v):return re.findall(r'1\.2\.643(?:\.\d+)+',text(v))
def sp(path,name,entity,expected):
    ws=load_workbook(path,read_only=True,data_only=True)['Детализация по СП'];rows=[]
    for r in ws.iter_rows(min_row=7,values_only=True):
        if text(r[0])!='Республика Татарстан' or not text(r[1]):continue
        planned=oids(r[4]);passed=oids(r[5]);failed=oids(r[6])
        if not planned:continue
        rows.append({'mo':text(r[1]),'moOid':text(r[2]),'unit':f'Здание {text(r[3])}','unitOid':', '.join(planned),'buildingIds':text(r[3]),'registered':bool(passed),'partial':bool(passed and failed),'count':0,'plannedSubunits':len(set(planned)),'registeredSubunits':len(set(planned)&set(passed))})
    fact=sum(x['registered'] for x in rows)
    if (len(rows),fact)!=expected:raise ValueError((name,len(rows),fact,expected))
    return {'name':name,'entity':entity,'plan':len(rows),'fact':fact,'date':'18.08.2026','rows':rows}
def buildings(path,sheet,name,entity,start,expected):
    ws=load_workbook(path,read_only=True,data_only=True)[sheet];seen={}
    for r in ws.iter_rows(min_row=start,values_only=True):
        if text(r[1])!='Республика Татарстан' or not text(r[2]):continue
        k=(text(r[3]),text(r[4]));seen[k]={'mo':text(r[2]),'moOid':text(r[3]),'unit':text(r[5]) or f'Здание {text(r[4])}','unitOid':text(r[4]),'buildingIds':text(r[4]),'registered':text(r[6]).lower()=='да','count':0}
    rows=list(seen.values());fact=sum(x['registered'] for x in rows)
    if (len(rows),fact)!=expected:raise ValueError((name,len(rows),fact,expected))
    return {'name':name,'entity':entity,'plan':len(rows),'fact':fact,'date':'18.08.2026','rows':rows}

stationary=one('Отчет_Доля_ТВСП_СЭМД_Эпикриз_в_стационаре*18.08*.xlsx')
ambulatory=one('Отчет_Доля_ТВСП_СЭМД_Эпикриз по законч*18.08*.xlsx')
lab=one('Отчет_Доля_ТВСП_СЭМД_Протокол лабораторного*18.08*.xlsx')
diag=one('Доля_ТВСП*диагностических*18.08*.xlsx')
smp=one('Отчет_СМП_ТВСП*18.08*.xlsx')
data={
 'tvspStationary':sp(stationary,'ТВСП, передающие выписные эпикризы','объект контроля с ТВСП',(276,273)),
 'tvspAmbulatory':sp(ambulatory,'Амбулаторные ТВСП, передающие эпикриз/талон и/или протокол консультации','объект контроля с ТВСП',(461,458)),
 'tvspLaboratory':sp(lab,'КДЛ, передающие протоколы лабораторных исследований','объект контроля КДЛ',(127,123)),
 'tvspDiagnostic':buildings(diag,'Факт передачи','ТВСП, передающие протоколы диагностических исследований','объект контроля',7,(304,297)),
 'smpFederal':buildings(smp,'Факт передачи','Станции и подстанции СМП, передающие карты вызова','станция / подстанция СМП',6,(69,67)),
}
(SITE/'app/unit-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
mo_path=SITE/'app/mo-data.json';detail_path=SITE/'app/mo-details.json';oid_path=SITE/'app/mo-detail-oids.json'
mo=json.loads(mo_path.read_text(encoding='utf-8'));details=json.loads(detail_path.read_text(encoding='utf-8'));detail_oids=json.loads(oid_path.read_text(encoding='utf-8'))
for metric,dataset in data.items():
    previous={str(r.get('oid') or ''):r.get('fact') for r in mo.get(metric,{}).get('rows',[]) if r.get('oid')}
    grouped={}
    for row in dataset['rows']:
        key=row['moOid'] or row['mo']
        item=grouped.setdefault(key,{'name':row['mo'],'oid':row['moOid'],'volume':0,'registered':0})
        item['volume']+=1;item['registered']+=int(row['registered'])
    rows=[];metric_details={};metric_oids={}
    for item in grouped.values():
        fact=item['registered']/item['volume']*100 if item['volume'] else 0
        old=previous.get(item['oid'])
        rows.append({'name':item['name'],'oid':item['oid'],'fact':fact,'count':item['registered'],'previous':old,'trend':fact-old if old is not None else None})
        detail={'volume':item['volume'],'registered':item['registered']}
        metric_details[item['name']]=detail
        if item['oid']:metric_oids[item['oid']]=detail
    mo[metric]={'name':dataset['name'],'plan':100,'unit':'%','date':'18.08.2026','period':'01.01–17.08.2026','rows':rows}
    details[metric]=metric_details;detail_oids[metric]=metric_oids
mo_path.write_text(json.dumps(mo,ensure_ascii=False,indent=2),encoding='utf-8')
detail_path.write_text(json.dumps(details,ensure_ascii=False,indent=2),encoding='utf-8')
oid_path.write_text(json.dumps(detail_oids,ensure_ascii=False,indent=2),encoding='utf-8')
print({k:(v['fact'],v['plan']) for k,v in data.items()})
