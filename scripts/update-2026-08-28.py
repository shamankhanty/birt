import json, glob, re, subprocess
from collections import Counter
import openpyxl

APP = "/workspace/sites/rt-health-digital-dashboard/app"
SRC = "/workspace/scratch/1f5740e14d75/audit_28_08_pjuDug"
HOSP = "/workspace/scratch/1f5740e14d75/upload/Отчет по госпитализациям 01.01.-27.08. (Выгрузка из ГИС ЭЗ РТ 27.08.26).xlsx"

def load(name):
    with open(f"{APP}/{name}", encoding="utf-8") as f: return json.load(f)
def save(name, data):
    with open(f"{APP}/{name}", "w", encoding="utf-8") as f: json.dump(data, f, ensure_ascii=False, indent=2)
def n(v):
    try: return int(float(v or 0))
    except: return 0
def key(name):
    name=re.sub(r'^(?:Филиал\s+)?(?:ГАУЗ|ГБУЗ|ГБУ|ФГБУ|ФГАОУВО|ФГАОУ ВО|АО|ООО)(?:\s+РТ)?\s*','',str(name),flags=re.I)
    name=re.sub(r'["«»]','',name); name=re.sub(r'\s*\([^)]*\)\s*$','',name)
    return re.sub(r'\s+',' ',name).strip().lower()

def previous_file(name):
    return json.loads(subprocess.check_output(['git','show',f'HEAD^:app/{name}'],text=True))

previous_mo=previous_file('mo-data.json')
previous_status=previous_file('organization-status.json')
def prior(metric, oid, name):
    rows=previous_mo.get(metric,{}).get('rows',[])
    matches=[r for r in rows if (oid and r.get('oid')==oid) or key(r.get('name',''))==key(name)]
    return matches[0].get('fact') if len(matches)==1 else None

mo=load('mo-data.json'); details=load('mo-details.json'); oids=load('mo-detail-oids.json'); op=load('operational-mo.json')

# ЕПГУ: производные отрицательные столбцы источника намеренно не используются.
p=glob.glob(SRC+'/ПР_2_*.xlsx')[0]; w=openpyxl.load_workbook(p,read_only=True,data_only=True); ws=w['Отчет']
rows=[r for r in ws.iter_rows(min_row=4,values_only=True) if r[2]]
for metric, col, title in [('egpu',5,'Доля заявлений о прикреплении на ЕПГУ'),('egpu2days',6,'Доля заявлений о прикреплении на ЕПГУ, рассмотренных за 2 рабочих дня')]:
    out=[]; det={}; oidmap={}
    for r in rows:
        volume=n(r[4]); registered=n(r[col]); fact=registered/volume*100 if volume else 0
        warning='Накопительный состав выгрузки изменён; динамика к 23.08 не рассчитывается.'
        out.append({'name':r[2],'oid':r[3],'fact':fact,'count':registered,'previous':None,'trend':None,'sourceWarning':warning})
        d={'volume':volume,'registered':registered}; det[key(r[2])]=d; oidmap[str(r[3])]=d
    mo[metric]={'name':title,'plan':100 if metric=='egpu' else 90,'unit':'%','date':'28.08.2026','period':'01.01–28.08.2026','note':'Динамика отключена: накопительный объём снизился с 32 799 до 29 635 заявлений.','rows':out}
    details[metric]=det; oids[metric]=oidmap

# СЭМД 228: по решению заказчика пустое значение числителя трактуется как 0.
p=glob.glob(SRC+'/Количество_СЭМД*.xlsx')[0]; w=openpyxl.load_workbook(p,read_only=True,data_only=True); ws=w.active
rows=[r for r in ws.iter_rows(min_row=5,values_only=True) if r[1]]; out=[]; det={}; oidmap={}
for r in rows:
    volume=n(r[3]); registered=n(r[4]); d={'volume':volume,'registered':registered}
    fact=registered/volume*100 if volume else 0
    out.append({'name':r[1],'oid':r[2],'fact':fact,'count':registered,'previous':None,'trend':None,'sourceWarning':'Состав знаменателя изменился; динамика к предыдущему срезу не рассчитывается.'})
    det[key(r[1])]=d; oidmap[str(r[2])]=d
mo['semd228']={'name':mo['semd228']['name'],'plan':95,'unit':'%','date':'27.08.2026','period':'01.01–27.08.2026','note':'Пустые значения СЭМД №228 приняты равными 0 по управленческому решению.','rows':out}; details['semd228']=det; oids['semd228']=oidmap

# Госпитализации: лист 3 содержит полный знаменатель случаев стационарной помощи.
w=openpyxl.load_workbook(HOSP,read_only=True,data_only=True); ws=w['Лист3']; rows=[r for r in ws.iter_rows(min_row=6,values_only=True) if r[1]]
out=[]; det={}; oidmap={}
for r in rows:
    volume=n(r[3]); registered=n(r[4])+n(r[5]); d={'volume':volume,'registered':registered}; fact=registered/volume*100 if volume else 0
    warn='Значение выше 100%: требуется проверка повторных/исправленных СЭМД.' if fact>100 else None
    prev=prior('hospital',r[2],r[1]); row={'name':r[1],'oid':r[2],'fact':fact,'count':registered,'previous':prev,'trend':None if prev is None else fact-prev}
    if warn: row['sourceWarning']=warn
    out.append(row); det[key(r[1])]=d; oidmap[str(r[2])]=d
mo['hospital']={'name':mo['hospital']['name'],'plan':95,'unit':'%','date':'27.08.2026','period':'01.01–27.08.2026','note':'Знаменатель — лист 3. Фактические значения свыше 100% показаны без искажения, рейтинг ограничивается 100 баллами.','rows':out}; details['hospital']=det; oids['hospital']=oidmap
op['hospitalCount']={'name':'Количество госпитализаций','plan':None,'unit':'','date':'27.08.2026','period':'01.01–27.08.2026','mode':'count','rows':[{'name':r[1],'oid':r[2],'fact':n(r[3]),'count':n(r[3]),'previous':None,'trend':None} for r in rows]}

# МАХ: слоты не являются частью расчёта; берём только два согласованных абсолютных поля листа 1.
p=glob.glob(SRC+'/ТМК_МАХ*.xlsx')[0]; w=openpyxl.load_workbook(p,read_only=True,data_only=True); ws=w['Лист1']; rows=[r for r in ws.iter_rows(min_row=3,values_only=True) if r[0]]
for metric,col,title in [('tmkMaxCount',3,'Количество проведённых ТМК посредством МАХ'),('elnMaxCount',4,'Количество ЛВН, закрытых после ТМК посредством МАХ')]:
    op[metric]={'name':title,'plan':10000,'unit':'','date':'28.08.2026','period':'01.08–28.08.2026','mode':'count','direction':'higher','note':'Слоты не используются в расчёте.','rows':[{'name':r[0],'fact':n(r[col]),'count':n(r[col]),'previous':None,'trend':None} for r in rows]}

# Категории отказов РЭМД.
p=glob.glob(SRC+'/Отчёт - Отказы*.xlsx')[0]; w=openpyxl.load_workbook(p,read_only=True,data_only=True); ws=w.active
c=Counter()
for r in ws.iter_rows(min_row=7,values_only=True):
    if r[2]: c[str(r[8] or r[7] or 'Без описания')]+=n(r[9])
errors=load('error-categories.json'); errors.update({'total':sum(c.values()),'period':'01.01–28.08.2026','items':[{'name':k,'count':v} for k,v in c.most_common()]})

# ТВСП: федеральный итог хранится отдельно от управленческого подсчёта всех СП внутри здания.
units=load('unit-data.json')
tvsp_specs=[
 ('tvspAmbulatory','Отчет_Доля_ТВСП_СЭМД_Эпикриз по законч*.xlsx','Амбулаторные ТВСП, передающие эпикриз/талон и/или протокол консультации','объект контроля с ТВСП'),
 ('tvspStationary','Отчет_Доля_ТВСП_СЭМД_Эпикриз_в_стационаре*.xlsx','ТВСП, передающие выписные эпикризы','объект контроля с ТВСП'),
 ('tvspLaboratory','Отчет_Доля_ТВСП_СЭМД_Протокол лабораторного*.xlsx','КДЛ, передающие протокол лабораторного исследования','КДЛ / лаборатория'),
]
def oid_list(v):
    if not v or str(v).strip().lower()=='нет': return []
    return re.findall(r'\d+(?:\.\d+)+',str(v))
for metric,pat,title,entity in tvsp_specs:
    p=glob.glob(SRC+'/'+pat)[0]; w=openpyxl.load_workbook(p,read_only=True,data_only=True)
    s=list(w['Сводный'].iter_rows(min_row=7,max_row=7,values_only=True))[0]; plan=n(s[2]); fact=n(s[3])
    raw=[]
    for r in w['Детализация по СП'].iter_rows(min_row=7,values_only=True):
        if r[0]!='Республика Татарстан' or not r[1]: continue
        planned=oid_list(r[4]); transmitted=oid_list(r[5]); registered=sorted(set(planned)&set(transmitted)); missing=sorted(set(planned)-set(registered))
        raw.append({'mo':r[1],'moOid':str(r[2] or ''),'unit':f'Здание {r[3]}','unitOid':', '.join(planned),'buildingIds':str(r[3] or ''),'registered':len(registered)>0,'partial':0<len(registered)<len(planned),'count':0,'plannedSubunits':len(planned),'registeredSubunits':len(registered)})
    units[metric]={'name':title,'entity':entity,'plan':plan,'fact':fact,'date':'28.08.2026','rows':raw}
    grouped={}
    for item in raw:
        g=grouped.setdefault(item['moOid'],{'name':item['mo'],'volume':0,'registered':0})
        g['volume']+=item['plannedSubunits']; g['registered']+=item['registeredSubunits']
    out=[]; det={}; oidmap={}
    for oid,g in grouped.items():
        d={'volume':g['volume'],'registered':g['registered']}; value=g['registered']/g['volume']*100 if g['volume'] else 0
        prev=prior(metric,oid,g['name']); row={'name':g['name'],'oid':oid,'fact':value,'count':g['registered'],'previous':prev,'trend':None if prev is None else value-prev}
        if metric=='tvspAmbulatory' and ('№ 18' in g['name'] or '№18' in g['name']): row['sourceWarning']='Применимость двух эндоскопических кабинетов требует уточнения.'
        out.append(row); det[key(g['name'])]=d; oidmap[oid]=d
    mo[metric]={'name':mo[metric]['name'],'plan':100,'unit':'%','date':'28.08.2026','period':'январь–август 2026','note':'Управленческий уровень рассчитан по всем обязательным СП; федеральная карточка — по объектам/зданиям.','rows':out}
    details[metric]=det; oids[metric]=oidmap

# Для двух объектных отчётов обновляем региональный итог и дату; детальная структура сохраняется.
units['tvspDiagnostic'].update({'plan':304,'fact':298,'date':'27.08.2026'})
units['smpFederal'].update({'plan':69,'fact':68,'date':'27.08.2026'})

# Свидетельства: обновляем реестр и управленческий остаток с трёхдневным сроком.
from datetime import datetime
def certificate(metric,pattern,start,name_col,date_col,status_col,title):
    p=glob.glob(SRC+'/'+pattern)[0]; w=openpyxl.load_workbook(p,read_only=True,data_only=True); ws=w.active
    raw=[r for r in ws.iter_rows(min_row=start,values_only=True) if r[name_col] and not str(r[name_col]).startswith('Где в столбце')]; grouped={}; cutoff=datetime(2026,8,25,23,59,59)
    for r in raw:
        g=grouped.setdefault(str(r[name_col]),{'volume':0,'registered':0,'overdueVolume':0,'overdueRegistered':0,'gracePending':0})
        registered=str(r[status_col] or '').strip()=='Зарегистрирован'; g['volume']+=1; g['registered']+=int(registered)
        if isinstance(r[date_col],datetime) and r[date_col]<=cutoff:
            g['overdueVolume']+=1; g['overdueRegistered']+=int(registered)
        elif not registered: g['gracePending']+=1
    old_by_key={key(r['name']):r for r in previous_mo.get(metric,{}).get('rows',[])}; out=[]; det={}
    for org,g in grouped.items():
        fact=g['registered']/g['volume']*100; old=old_by_key.get(key(org)); prev=old.get('fact') if old else None
        attention=g['overdueRegistered']/g['overdueVolume']*100 if g['overdueVolume'] else 100
        row={'name':org,'fact':fact,'count':g['registered'],'attentionFact':attention,'overdueVolume':g['overdueVolume'],'overdueRegistered':g['overdueRegistered'],'gracePending':g['gracePending'],'previous':prev,'trend':None if prev is None else fact-prev}
        if old and old.get('oid'): row['oid']=old['oid']
        out.append(row); det[key(org)]=g
    mo[metric]={'name':title,'plan':99,'unit':'%','date':'28.08.2026','period':'01.01–28.08.2026','note':'В управленческий остаток не включены свидетельства последних трёх календарных дней.','rows':out}; details[metric]=det

certificate('birth','Свидетельства_о_рождении*.xlsx',3,0,2,3,mo['birth']['name'])
certificate('death','I_Свид-ва*.xlsx',4,1,12,93,mo['death']['name'])

# ЭЛМК: обновляем только утверждённый перечень 74 МО по OID, нулевой факт остаётся нейтральным до подтверждения применимости.
status=load('organization-status.json'); p=glob.glob(SRC+'/Медкнижки*.xlsx')[0]; w=openpyxl.load_workbook(p,read_only=True,data_only=True); ws=w['Отчет РЭМД по МО']
counts={str(r[2]):n(r[125]) for r in ws.iter_rows(min_row=8,values_only=True) if r[2]}
old_rows=previous_status['elmk']['rows']; elmk_rows=[]
for old in old_rows:
    count=counts.get(str(old.get('oid')),0); fact=100 if count>0 else 0; prev=old.get('fact')
    elmk_rows.append({**old,'count':count,'fact':fact,'previous':prev,'trend':None if prev is None else fact-prev})
status['elmk'].update({'date':'28.08.2026','note':f'Плановый перечень — 74 МО. На 28.08.2026 передача подтверждена у {sum(r["fact"]>0 for r in elmk_rows)} МО; ноль не трактуется как нарушение без проверки лицензии.','rows':elmk_rows})

# Матрица применимости: наличие знаменателя/планового объекта в профильном источнике подтверждает применимость.
registry=load('mo-registry.json'); by_oid={x['oid']:x for x in registry['organizations']}
for metric in ['egpu','egpu2days','semd228','hospital','ambulatoryCase','tvspStationary','tvspAmbulatory','tvspLaboratory','tvspDiagnostic','smpFederal']:
    for row in mo.get(metric,{}).get('rows',[]):
        org=by_oid.get(str(row.get('oid',''))); d=oids.get(metric,{}).get(str(row.get('oid','')))
        if org and d and d.get('volume',0)>0 and metric not in org['applicable']: org['applicable'].append(metric)
for org in registry['organizations']: org['applicable']=sorted(set(org['applicable']))

save('mo-data.json',mo); save('mo-details.json',details); save('mo-detail-oids.json',oids); save('operational-mo.json',op); save('error-categories.json',errors); save('unit-data.json',units); save('organization-status.json',status); save('mo-registry.json',registry)
