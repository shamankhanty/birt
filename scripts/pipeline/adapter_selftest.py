#!/usr/bin/env python3
from __future__ import annotations
import hashlib,json,shutil,tempfile
from datetime import date
from pathlib import Path
import openpyxl
from openpyxl import Workbook

ROOT=Path(__file__).resolve().parents[2]
import sys
sys.path.insert(0,str(ROOT/'scripts/pipeline'))
from adapters import (adapt_egpu,adapt_hospital,adapt_certificates,adapt_max,adapt_errors,adapt_physicians,adapt_preventive,adapt_waybill,
 adapt_tvsp_subunits,adapt_tvsp_buildings,adapt_presence,adapt_short_input,adapt_fap,adapt_asu_smp)
from staging_runner import stage

def sha_tree(p):
 return {x.relative_to(p).as_posix():hashlib.sha256(x.read_bytes()).hexdigest() for x in p.rglob('*') if x.is_file()}

def wb_save(path,sheets):
 wb=Workbook(); wb.remove(wb.active)
 for name,writer in sheets:
  ws=wb.create_sheet(name); writer(ws)
 wb.save(path)

def oid_name():
 reg=json.loads((ROOT/'app/mo-registry.json').read_text(encoding='utf-8'))['organizations'][0]
 return reg['oid'],reg['name']

def build_sources(d:Path):
 oid,name=oid_name(); end='30.09.2026'
 p=d/f'ПР_2_{end}.xlsx'
 def eg(ws): ws.cell(4,3,name);ws.cell(4,4,oid);ws.cell(4,5,100);ws.cell(4,6,98);ws.cell(4,7,90)
 wb_save(p,[('Отчет',eg)])
 h=d/f'Отчет_по_госпитализациям_{end}.xlsx'
 def hosp(ws):ws.cell(6,2,name);ws.cell(6,3,oid);ws.cell(6,4,100);ws.cell(6,5,40);ws.cell(6,6,35)
 wb_save(h,[('Лист3',hosp)])
 b=d/f'Свидетельства_о_рождении_01.01-30.09.2026.xlsx'
 def cert(ws):ws.cell(3,1,name);ws.cell(3,2,'B1');ws.cell(3,4,'Зарегистрирован');ws.cell(4,1,name);ws.cell(4,2,'B2');ws.cell(4,4,'Создан')
 wb_save(b,[('Лист1',cert)])
 de=d/f'I_Свид-ва о смерти_01.01-30.09.2026.xlsx';wb_save(de,[('Лист1',cert)])
 mx=d/f'ТМК_МАХ_{end}.xlsx'
 def maxw(ws):ws.cell(3,1,name);ws.cell(3,4,12);ws.cell(3,5,7)
 wb_save(mx,[('Лист1',maxw)])
 er=d/f'Отчёт - Отказы_{end}.xlsx'
 def erw(ws):
  ws.cell(7,2,oid);ws.cell(7,3,name);ws.cell(7,9,'Ошибка тест');ws.cell(7,10,15)
  ws.cell(8,9,'Ошибка без МО');ws.cell(8,10,5)
 wb_save(er,[('Лист1',erw)])
 rem=d/f'1.Отчет СЭМД_РЭМД_{end}.xlsx'
 def remw(ws):ws.cell(8,2,name);ws.cell(8,3,oid);ws.cell(8,81,80);ws.cell(8,108,75)
 wb_save(rem,[('Отчет РЭМД по МО',remw)])
 fo=d/f'Количество_СЭМД_Результаты_{end}.xlsx'
 def fow(ws):ws.cell(5,2,name);ws.cell(5,3,oid);ws.cell(5,4,100)
 wb_save(fo,[('Лист1',fow)])
 ph=d/f'Отчёт_по_врачам_{end}.xlsx'
 specs=['Акушер-гинеколог','Врач общей практики','Кардиолог','Онколог','Офтальмолог','Педиатр','Стоматолог','Терапевт','Хирург']
 def allw(ws):ws.cell(9,1,'Республика Татарстан');ws.cell(9,2,oid);ws.cell(9,3,name);ws.cell(9,4,'III уровень');ws.cell(9,5,10);ws.cell(9,7,8)
 def spw(ws):
  for i,sp in enumerate(specs,9):
   ws.cell(i,1,'Республика Татарстан');ws.cell(i,2,oid);ws.cell(i,3,name);ws.cell(i,5,sp);ws.cell(i,6,10);ws.cell(i,7,9);ws.cell(i,8,2);ws.cell(i,9,1);ws.cell(i,10,6);ws.cell(i,14,0.6)
 def subj(ws):
  for i,sp in enumerate(specs,10):ws.cell(i,1,'Республика Татарстан');ws.cell(i,2,sp);ws.cell(i,3,10);ws.cell(i,7,6)
 wb_save(ph,[('Все врачи_Детализация по МО',allw),('Врачи по спец-тям_По МО',spw),('Врачи по спец-тям_По субъекту',subj)])
 way=d/f'Отчёт_по_использованию_системы_24.09.2026_30.09.2026.xlsx'
 def wayw(ws):ws['B2']='тест';vals=[1,name,10,4,6,8,5,1,1,12,5,7,4];
 def fillway(ws):
  ws['B2']='тест'; vals=[1,name,10,4,6,8,5,1,1,12,5,7,4]
  for j,v in enumerate(vals,1):ws.cell(14,j,v)
 wb_save(way,[('Лист1',fillway)])
 def spfile(filename):
  path=d/filename
  def summary(ws): ws.cell(7,3,1);ws.cell(7,4,1)
  def detail(ws): ws.cell(7,1,'Республика Татарстан');ws.cell(7,2,name);ws.cell(7,3,oid);ws.cell(7,4,'23575');ws.cell(7,5,oid+'.0.1');ws.cell(7,6,oid+'.0.1')
  wb_save(path,[('Сводный',summary),('Детализация по СП',detail)]);return path
 tvspa=spfile(f'Отчет_Доля_ТВСП_СЭМД_Эпикриз по законченному_{end}.xlsx')
 tvsps=spfile(f'Отчет_Доля_ТВСП_СЭМД_Эпикриз_в_стационаре_{end}.xlsx')
 tvspl=spfile(f'Отчет_Доля_ТВСП_СЭМД_Протокол лабораторного_{end}.xlsx')
 def building(filename,rownum):
  path=d/filename
  def wr(ws):ws.cell(rownum,2,'Республика Татарстан');ws.cell(rownum,3,name);ws.cell(rownum,4,oid);ws.cell(rownum,5,'23575');ws.cell(rownum,6,'Тестовый объект');ws.cell(rownum,7,'Да')
  wb_save(path,[('Факт передачи',wr)]);return path
 diag=building(f'Доля_ТВСП_диагностических_{end}.xlsx',8); smpt=building(f'Отчет_СМП_ТВСП_{end}.xlsx',7)
 tmk=d/f'Отчет_ТМК_РЭМД_{end}.xlsx'
 def tmkw(ws):ws.cell(7,2,'Республика Татарстан');ws.cell(7,3,name);ws.cell(7,4,oid);ws.cell(7,5,5)
 wb_save(tmk,[('Детализированный отчет',tmkw)])
 elmk=d/f'Медкнижки_{end}.xlsx'
 def elmkw(ws):ws.cell(8,2,name);ws.cell(8,3,oid);ws.cell(8,126,9)
 wb_save(elmk,[('Отчет РЭМД по МО',elmkw)])
 short=d/f'Случаи краткого ввода_{end}.xlsx'
 def shortw(ws):ws.cell(7,1,name);ws.cell(7,2,oid);ws.cell(7,4,2);ws.cell(7,5,1);ws.cell(7,6,1);ws.cell(7,7,3)
 wb_save(short,[('Краткий ввод',shortw)])
 fap=d/f'Отчет_по_ФАП_и_ФП_{end}.xlsx'
 def fapw(ws):ws.cell(5,1,name);ws.cell(5,4,12);ws.cell(6,1,name);ws.cell(6,4,0)
 wb_save(fap,[('Лист1',fapw)])
 asu=d/f'АСУ_ССМП_{end}.xlsx'
 wb=Workbook();ws=wb.active;ws.cell(9,1,name);ws.cell(9,1).font=openpyxl.styles.Font(bold=True);ws.cell(9,2,100);ws.cell(9,11,92);wb.save(asu)
 return {'egpu':p,'hospital':h,'birth':b,'death':de,'max':mx,'errors':er,'remd':rem,'foms':fo,'physicians':ph,'waybill':way,
 'tvspa':tvspa,'tvsps':tvsps,'tvspl':tvspl,'diag':diag,'smpt':smpt,'tmk':tmk,'elmk':elmk,'short':short,'fap':fap,'asu':asu}

def fresh(base:Path,label):
 p=base/label;shutil.copytree(ROOT/'app',p);return p

def main():
 canonical=sha_tree(ROOT/'app'); results=[]
 with tempfile.TemporaryDirectory() as td:
  d=Path(td);src=build_sources(d);end=date(2026,9,30);test_oid,_=oid_name()
  cases=[
   ('egpu',lambda a:adapt_egpu(a,src['egpu'],end)),('hospital',lambda a:adapt_hospital(a,src['hospital'],end)),
   ('birth',lambda a:adapt_certificates(a,src['birth'],end,'birth','birth_certificates')),('death',lambda a:adapt_certificates(a,src['death'],end,'death','death_certificates')),
   ('max',lambda a:adapt_max(a,src['max'],end)),('errors',lambda a:adapt_errors(a,src['errors'],end)),
   ('physicians',lambda a:adapt_physicians(a,src['physicians'],end,ROOT/'app/mo-registry.json')),
   ('preventive',lambda a:adapt_preventive(a,src['remd'],src['foms'],end,ROOT/'app/mo-registry.json')),
   ('waybill',lambda a:adapt_waybill(a,src['waybill'],end)),
   ('tvsp-ambulatory',lambda a:adapt_tvsp_subunits(a,src['tvspa'],end,'tvspAmbulatory','tvsp_ambulatory','Амбулаторные ТВСП, передающие эпикриз/талон и/или протокол консультации','объект контроля с ТВСП')),
   ('tvsp-stationary',lambda a:adapt_tvsp_subunits(a,src['tvsps'],end,'tvspStationary','tvsp_stationary','ТВСП, передающие выписные эпикризы','объект контроля с ТВСП')),
   ('tvsp-laboratory',lambda a:adapt_tvsp_subunits(a,src['tvspl'],end,'tvspLaboratory','tvsp_laboratory','КДЛ, передающие протокол лабораторного исследования','КДЛ / лаборатория')),
   ('tvsp-diagnostic',lambda a:adapt_tvsp_buildings(a,src['diag'],end,'tvspDiagnostic','tvsp_diagnostic','ТВСП, передающие протоколы диагностических исследований','объект контроля',8)),
   ('smp-tvsp',lambda a:adapt_tvsp_buildings(a,src['smpt'],end,'smpFederal','smp_tvsp','Станции и подстанции СМП, передающие карты вызова','станция / подстанция СМП',7)),
   ('tmk-remd',lambda a:adapt_presence(a,src['tmk'],end,'tmk_remd')),
   ('elmk',lambda a:adapt_presence(a,src['elmk'],end,'elmk')),
   ('short-input',lambda a:adapt_short_input(a,src['short'],end)),
   ('fap',lambda a:adapt_fap(a,src['fap'],end)),
   ('asu-smp',lambda a:adapt_asu_smp(a,src['asu'],end)),
  ]
  for label,fn in cases:
   a=fresh(d,'app-'+label);r=fn(a);assert r.status=='PASS',(label,r);assert r.changedFiles
   if label=='errors':
    payload=json.loads((a/'error-categories.json').read_text(encoding='utf-8')); bd=payload['organizationBreakdown']
    assert payload['total']==20,payload['total'];assert bd['attributedErrors']==15,bd;assert bd['unassignedErrors']==5,bd;assert bd['coveragePercent']==75.0,bd
    assert bd['organizations'][0]['oid']==test_oid,bd['organizations'][0];assert bd['organizations'][0]['topCategories'][0]['name']=='Ошибка тест'
   results.append({'adapter':label,'changed':r.changedFiles,'facts':r.facts})
  # End-to-end staging can be formally PASS on a non-rating operational source.
  inp=d/'only-errors';inp.mkdir();shutil.copy2(src['errors'],inp/src['errors'].name);out=d/'candidate';m=stage(inp,out);assert m['status']=='PASS',m;assert m['formalValidation']['status']=='PASS',m['formalValidation'];assert m['changedFiles']==['error-categories.json'],m['changedFiles']
 assert sha_tree(ROOT/'app')==canonical,'canonical app changed'
 print(json.dumps({'status':'PASS','adapters':len(results),'adapterResults':results,'canonicalUnchanged':True,'endToEndStaging':'PASS'},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
