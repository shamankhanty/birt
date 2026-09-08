#!/usr/bin/env python3
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = Path('/workspace/scratch/1f5740e14d75/analysis_jan_aug/Закрытые_случаи_ДОГВН_ПМО_и_зарегистрированные_СЭМД_122_и_228_ Январь - Август.xlsx')

def load(name):
    return json.loads((ROOT/name).read_text(encoding='utf-8'))

def save(name, data):
    (ROOT/name).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

# Full August physician datasets: expose them in the monthly MO mode.
phys = load('app/physician-metrics.json')
monthly = load('app/monthly-mo.json')
for metric_id, dataset in phys['datasets'].items():
    monthly[metric_id] = {
        'unit': dataset['unit'],
        'previousLabel': 'Июль',
        'currentLabel': 'Август',
        'rows': [{
            'name': row['name'], 'june': None, 'july': row['fact'], 'change': None,
            'juneQuantity': None,
            'julyQuantity': f"{row['count']:,} / {row['volume']:,}".replace(',', ' '),
        } for row in dataset['rows']],
    }
save('app/monthly-mo.json', monthly)

# Preventive examinations: adults use closed cases (D+E); child-only rows use
# registered SЭМД plus the workbook's explicit "without SЭМД" balance.
ws = load_workbook(SRC, read_only=True, data_only=True).active
rows=[]
for values in ws.iter_rows(min_row=5, values_only=True):
    if not values[1] or str(values[1]).strip().casefold().startswith('итого'):
        continue
    name, oid = str(values[1]).strip(), str(values[2]).strip()
    adult = isinstance(values[3], (int,float)) or isinstance(values[4], (int,float))
    semd122, semd228 = int(values[9] or 0), int(values[13] or 0)
    selected=max(semd122,semd228)
    foms=(int(values[3] or 0)+int(values[4] or 0)) if adult else selected+int(values[5] or 0)
    rows.append({'name':name,'oid':oid,'child':not adult,'semd122':semd122,'semd228':semd228,
                 'selected':selected,'selectedType':'122' if semd122>=semd228 else '228',
                 'foms':foms,'share':selected/foms*100 if foms else None,
                 'oldShare':None,'change':None,'issues':[] if foms else ['Знаменатель равен 0']})
num=sum(r['selected'] for r in rows); den=sum(r['foms'] for r in rows)
audit={'summary':{'status':'ready','year':2026,'formula':'MAX(СЭМД 122; СЭМД 228)',
    'period122':'01.01.2026–31.08.2026','period228':'01.01.2026–31.08.2026',
    'source122':SRC.name,'source228':SRC.name,'organizations':len(rows),'changed':0,
    'changedChildren':0,'over100':0,'zeroDenominatorWithSemd':0,
    'missing':sum(bool(r['issues']) for r in rows),'childrenMissing':False,
    'childOrganizations':sum(r['child'] for r in rows),'numerator':num,'denominator':den,
    'share':num/den*100,'selected122':sum(r['selectedType']=='122' for r in rows),
    'selected228':sum(r['selectedType']=='228' for r in rows),
    'selectedEqual':sum(r['semd122']==r['semd228'] for r in rows),
    'note':'В расчёт включены взрослые и детские МО. Для детских строк знаменатель восстановлен как зарегистрированные СЭМД плюс случаи без СЭМД.'},
    'rows':rows}
save('app/preventive-semd-audit.json', audit)

mo=load('app/mo-data.json'); details=load('app/mo-details.json'); oids=load('app/mo-detail-oids.json')
mo['semd228']['rows']=[{'name':r['name'],'fact':r['share'],'previous':None,'trend':None} for r in rows]
details['semd228']={r['name'].lower().replace('гауз ', '').replace('"',''):{'volume':r['foms'],'registered':r['selected']} for r in rows}
oids['semd228']={r['oid']:{'volume':r['foms'],'registered':r['selected']} for r in rows if r['oid']}
save('app/mo-data.json',mo); save('app/mo-details.json',details); save('app/mo-detail-oids.json',oids)
monthly=load('app/monthly-mo.json')
monthly['semd228']={'unit':'%','previousLabel':'Июль','currentLabel':'Август','rows':[
    {'name':r['name'],'june':None,'july':r['share'],'change':None,'juneQuantity':None,
     'julyQuantity':f"{r['selected']:,} / {r['foms']:,}".replace(',', ' ')} for r in rows]}
save('app/monthly-mo.json',monthly)
print(json.dumps({'physicianDatasets':len(phys['datasets']),'preventiveOrganizations':len(rows),
                  'preventiveChildren':sum(r['child'] for r in rows),'preventiveNumerator':num,
                  'preventiveDenominator':den,'preventiveShare':num/den*100},ensure_ascii=False,indent=2))
