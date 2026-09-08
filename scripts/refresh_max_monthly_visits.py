#!/usr/bin/env python3
"""Add the approved 31.07/31.08 cumulative visit snapshots to monthly runtime."""
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = Path("/workspace/scratch/e3dedb97a4e9/source-audit")

def source(folder: str) -> Path:
    return next((SOURCE_ROOT / folder).rglob("ТМК_МАХ*.xlsx"))

def rows(path: Path) -> dict[str, int]:
    ws = load_workbook(path, read_only=True, data_only=True)["Лист4"]
    return {
        str(row[0]).strip(): int(row[5] or 0)
        for row in ws.iter_rows(min_row=3, values_only=True)
        if row[0]
    }

july = rows(source("july"))
august = rows(source("august"))
names = sorted(set(july) | set(august))
payload_path = ROOT / "app" / "max-monthly-visits.json"
payload = {
    "name": "Количество записей к врачу посредством МАХ",
    "plan": None,
    "unit": "",
    "previousLabel": "На 31.07",
    "currentLabel": "На 31.08",
    "rows": [
        {
            "name": name,
            "june": july.get(name),
            "july": august.get(name),
            "change": august[name] - july[name] if name in july and name in august else None,
            "juneQuantity": str(july[name]) if name in july else None,
            "julyQuantity": str(august[name]) if name in august else None,
            **({"sourceWarning": "Нет строки в выгрузке на 31.07.2026"} if name not in july else {}),
            **({"sourceWarning": "Нет строки в выгрузке на 31.08.2026"} if name not in august else {}),
        }
        for name in names
    ],
}
payload_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"julyRows": len(july), "augustRows": len(august)}, ensure_ascii=False))
