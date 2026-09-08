#!/usr/bin/env python3
"""Build the dashboard dataset from the July electronic-waybill report."""

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/workspace/scratch/8e713f374d53/upload/Отчёт_по_использованию_системы_01_07_2026.xlsx")
OUTPUT = ROOT / "app" / "electronic-waybill.json"


def number(value):
    return int(value or 0)


ws = load_workbook(SOURCE, read_only=True, data_only=True)["Статистика"]

rows = []
for values in ws.iter_rows(min_row=15, values_only=True):
    if not isinstance(values[0], (int, float)) or not values[1]:
        continue
    vehicles = number(values[2])
    moved = number(values[5])
    rows.append(
        {
            "sourceNumber": number(values[0]),
            "name": str(values[1]).strip(),
            "vehicles": vehicles,
            "ambulanceVehicles": number(values[3]),
            "otherVehicles": number(values[4]),
            "moved": moved,
            "movementShare": moved / vehicles * 100 if vehicles else None,
            "drivers": number(values[6]),
            "mechanics": number(values[7]),
            "medics": number(values[8]),
            "waybills": number(values[9]),
            "ambulanceWaybills": number(values[10]),
            "otherWaybills": number(values[11]),
            "driversWithWaybills": number(values[12]),
        }
    )

summary = {
    "organizations": number(ws["C7"].value),
    "vehicles": number(ws["C8"].value),
    "vehiclesWithWaybills": number(ws["E8"].value),
    "vehiclesWithWaybillsShare": float(ws["G8"].value or 0) * 100,
    "vehiclesWithMovement": number(ws["H8"].value),
    "vehiclesWithMovementShare": float(ws["I8"].value or 0) * 100,
    "ambulanceVehicles": number(ws["C9"].value),
    "ambulanceVehiclesWithWaybills": number(ws["E9"].value),
    "ambulanceVehiclesWithWaybillsShare": float(ws["G9"].value or 0) * 100,
    "ambulanceVehiclesWithMovement": number(ws["H9"].value),
    "ambulanceVehiclesWithMovementShare": float(ws["I9"].value or 0) * 100,
    "kazanAmbulanceVehicles": number(ws["C10"].value),
    "kazanAmbulanceVehiclesWithWaybills": number(ws["E10"].value),
    "kazanAmbulanceVehiclesWithWaybillsShare": float(ws["G10"].value or 0) * 100,
    "kazanAmbulanceVehiclesWithMovement": number(ws["H10"].value),
    "kazanAmbulanceVehiclesWithMovementShare": float(ws["I10"].value or 0) * 100,
    "otherVehicles": number(ws["C11"].value),
    "otherVehiclesWithWaybills": number(ws["E11"].value),
    "otherVehiclesWithWaybillsShare": float(ws["G11"].value or 0) * 100,
    "otherVehiclesWithMovement": number(ws["H11"].value),
    "otherVehiclesWithMovementShare": float(ws["I11"].value or 0) * 100,
}

detail = {
    "organizations": len(rows),
    "vehicles": sum(row["vehicles"] for row in rows),
    "vehiclesWithMovement": sum(row["moved"] for row in rows),
    "waybills": sum(row["waybills"] for row in rows),
    "driversWithWaybills": sum(row["driversWithWaybills"] for row in rows),
    "zeroMovementOrganizations": sum(row["vehicles"] > 0 and row["moved"] == 0 for row in rows),
    "zeroVehicleOrganizations": sum(row["vehicles"] == 0 for row in rows),
}

quality = [
    {
        "field": "Транспортные средства в системе",
        "summary": summary["vehicles"],
        "detail": detail["vehicles"],
        "difference": detail["vehicles"] - summary["vehicles"],
        "note": "Итоговый блок не воспроизводится суммой строк организаций.",
    },
    {
        "field": "ТС с движением более 1 км",
        "summary": summary["vehiclesWithMovement"],
        "detail": detail["vehiclesWithMovement"],
        "difference": detail["vehiclesWithMovement"] - summary["vehiclesWithMovement"],
        "note": "Для показателя по РТ используется итоговый блок; рейтинг МО построен только по детализации.",
    },
    {
        "field": "ТС, на которые создавались путевые листы",
        "summary": summary["vehiclesWithWaybills"],
        "detail": summary["ambulanceVehiclesWithWaybills"] + summary["otherVehiclesWithWaybills"],
        "difference": summary["ambulanceVehiclesWithWaybills"] + summary["otherVehiclesWithWaybills"] - summary["vehiclesWithWaybills"],
        "note": "Общий итог не равен сумме автомобилей СМП и остальных автомобилей.",
    },
]

payload = {
    "title": "Электронный путевой лист",
    "source": SOURCE.name,
    "period": "01.07–31.07.2026",
    "formed": None,
    "plan": None,
    "summary": summary,
    "detail": detail,
    "quality": quality,
    "rows": rows,
}

OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"output": str(OUTPUT), "summary": summary, "detail": detail}, ensure_ascii=False))
