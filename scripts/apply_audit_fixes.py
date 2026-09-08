import json
import re
from pathlib import Path

from openpyxl import load_workbook


SITE = Path(__file__).resolve().parents[1]
SOURCE = Path("/workspace/scratch/8e713f374d53/upload/2 слайд ПР_2_Кол-во заявл на прикр. от ЕГПУ за период. ИП ООГУЗ. Обраб. 2 суток. 01.08.-05.08. (Выгрузка из гИС ЭЗ РТ 05.08.26).xlsx")


def source_key(name: str) -> str:
    value = re.sub(
        r"^(?:Филиал\s+)?(?:ГАУЗ|ГБУЗ|ГБУ|ФГБУ|ФГАОУВО|ФГАОУ ВО|АО|ООО)(?:\s+РТ)?\s*",
        "",
        str(name),
        flags=re.I,
    )
    value = re.sub(r'["«»]', "", value)
    value = re.sub(r"\s*\([^)]*\)\s*$", "", value)
    return re.sub(r"\s+", " ", value).strip().lower()


def read_period(ws):
    rows = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        name, oid, submitted, final, two_days = row[:5]
        if not name or not isinstance(submitted, (int, float)):
            continue
        rows[str(oid)] = {
            "name": str(name),
            "submitted": submitted,
            "final": final or 0,
            "two_days": two_days or 0,
        }
    return rows


mo_path = SITE / "app" / "mo-data.json"
details_path = SITE / "app" / "mo-details.json"
mo = json.loads(mo_path.read_text(encoding="utf-8"))
details = json.loads(details_path.read_text(encoding="utf-8"))

wb = load_workbook(SOURCE, read_only=True, data_only=True)
current = read_period(wb["01.01.-05.08."])
previous = read_period(wb["01.01.-29.07."])

for metric, numerator in (("egpu", "final"), ("egpu2days", "two_days")):
    metric_rows = []
    metric_details = {}
    for oid, item in current.items():
        fact = item[numerator] / item["submitted"] * 100 if item["submitted"] else 0
        old = previous.get(oid)
        old_fact = (
            old[numerator] / old["submitted"] * 100
            if old and old["submitted"]
            else None
        )
        metric_rows.append(
            {
                "name": item["name"],
                "fact": fact,
                "previous": old_fact,
                "trend": fact - old_fact if old_fact is not None else None,
            }
        )
        metric_details[source_key(item["name"])] = {
            "volume": item["submitted"],
            "registered": item[numerator],
        }
    mo[metric]["rows"] = metric_rows
    details[metric] = metric_details

# Технические итоги не являются медицинскими организациями.
for dataset in mo.values():
    if isinstance(dataset, dict) and isinstance(dataset.get("rows"), list):
        dataset["rows"] = [
            row
            for row in dataset["rows"]
            if not re.match(r"^\s*(?:итого|всего по)", str(row.get("name", "")), re.I)
        ]

mo_path.write_text(json.dumps(mo, ensure_ascii=False, indent=2), encoding="utf-8")
details_path.write_text(json.dumps(details, ensure_ascii=False, indent=2), encoding="utf-8")
